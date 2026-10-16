#!/usr/bin/env python3
"""One-shot runner for marketing-ws03 /generate-html-for-user-stories (operator run)."""
from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
SCRIPTS = ROOT / "scripts"
PIPE = ROOT / "artifacts" / "evaluation" / ".pipeline"
INTAKE_PIPE = ROOT / "artifacts" / "intake" / ".pipeline"

if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))
from generate_user_story_source_analysis import (  # noqa: E402
    NO_EVIDENCE_EXCERPT,
    has_supporting_refs,
    match_type_from_evidence,
)

RULEBOOK_COLS = [
    "State",
    "ID",
    "Work Item Type",
    "Title 1",
    "Title 2",
    "Title 3",
    "Assigned To",
    "Area Path",
    "Description",
    "Acceptance Criteria",
    "Annexure ID1",
    "Fit Gap Status",
    "User Story Type",
]

TOKEN_RE = re.compile(r"[a-z0-9]+")
STOP = {
    "a",
    "an",
    "the",
    "to",
    "for",
    "and",
    "or",
    "in",
    "on",
    "of",
    "as",
    "is",
    "be",
    "that",
    "this",
    "with",
    "from",
    "need",
    "user",
    "story",
}


def _tokens(text: str) -> set[str]:
    return {t for t in TOKEN_RE.findall((text or "").lower()) if t not in STOP and len(t) > 2}


def _jaccard(a: set[str], b: set[str]) -> float:
    if not a or not b:
        return 0.0
    return len(a & b) / len(a | b)


def _run(cmd: list[str]) -> None:
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    if proc.returncode != 0:
        print(proc.stdout, proc.stderr, file=sys.stderr)
        raise SystemExit(proc.returncode)


def _xlsx_to_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    rows: list[dict[str, str]] = []
    for r in range(2, ws.max_row + 1):
        row = {headers[i - 1]: str(ws.cell(r, i).value or "").strip() for i in range(1, len(headers) + 1)}
        if not any(row.values()):
            continue
        rows.append(row)
    return headers, rows


def _legacy_manual_to_rulebook(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for row in rows:
        mapped = {col: "" for col in RULEBOOK_COLS}
        mapped["State"] = row.get("State", "")
        mapped["ID"] = row.get("ID", "")
        mapped["Work Item Type"] = row.get("Work Item Type", "")
        mapped["Title 1"] = row.get("Title 1", "")
        mapped["Title 2"] = row.get("Title 2", "")
        mapped["Title 3"] = row.get("Title 3", "")
        mapped["Area Path"] = row.get("Area Level 2", row.get("Area Path", ""))
        mapped["Description"] = row.get("Description", "")
        mapped["Acceptance Criteria"] = row.get("Acceptance Criteria", "")
        out.append(mapped)
    return out


def _rulebook_rows_from_xlsx(path: Path, legacy_manual: bool = False) -> list[dict[str, str]]:
    headers, rows = _xlsx_to_rows(path)
    if legacy_manual:
        return _legacy_manual_to_rulebook(rows)
    out: list[dict[str, str]] = []
    for row in rows:
        mapped = {col: row.get(col, "") for col in RULEBOOK_COLS}
        out.append(mapped)
    return out


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=RULEBOOK_COLS)
        writer.writeheader()
        for row in rows:
            writer.writerow({col: row.get(col, "") for col in RULEBOOK_COLS})


def _build_unified_kb(kb_index: dict, prereq_path: Path) -> dict:
    prereq = json.loads(prereq_path.read_text(encoding="utf-8"))
    requirements: list[dict] = []
    for bucket, source_type in (
        ("in_scope_objectives", "Pre-read Materials"),
        ("in_scope_modules", "Pre-read Materials"),
        ("out_of_scope", "Pre-read Materials"),
    ):
        for item in (prereq.get("scope") or {}).get(bucket) or []:
            requirements.append(
                {
                    "id": item.get("id"),
                    "statement": item.get("statement"),
                    "source_type": source_type,
                    "source_path": item.get("source_path"),
                    "source_locator": item.get("source_locator"),
                    "confidence": item.get("confidence", "explicit"),
                }
            )
    for chunk in kb_index.get("source_corpus_chunks") or []:
        text = (chunk.get("text") or "").strip()
        if len(text) < 40:
            continue
        requirements.append(
            {
                "id": f"CHUNK-{chunk.get('source_type')}-{chunk.get('locator')}",
                "statement": text[:500],
                "source_type": chunk.get("source_type"),
                "source_path": chunk.get("source_path"),
                "source_locator": chunk.get("locator"),
                "confidence": "explicit",
            }
        )
    features = []
    for feat in kb_index.get("feature_list_rows") or []:
        features.append(
            {
                "feature_name": feat.get("feature_name"),
                "workshop_code": feat.get("workshop_code"),
                "source_path": feat.get("source_path"),
                "source_locator": feat.get("source_locator"),
            }
        )
    return {
        "schema_version": "1.0",
        "kind": "unified-source-knowledge-base",
        "sources": [
            {"path": p, "role": "transcript"}
            for p in kb_index.get("inputs", {}).get("transcript_paths", [])
        ]
        + [{"path": p, "role": "mom"} for p in kb_index.get("inputs", {}).get("mom_paths", [])]
        + [{"path": p, "role": "preread"} for p in kb_index.get("inputs", {}).get("preread_paths", [])]
        + [{"path": kb_index.get("inputs", {}).get("sow_json"), "role": "feature_list"}],
        "requirements": requirements,
        "features": features,
        "prerequisite_kb_ref": str(prereq_path),
    }


def _best_matches(story_blob: str, unified_kb: dict, top: int = 3) -> list[tuple[float, dict]]:
    st = _tokens(story_blob)
    scored: list[tuple[float, dict]] = []
    for req in unified_kb.get("requirements") or []:
        score = _jaccard(st, _tokens(req.get("statement", "")))
        if score > 0:
            scored.append((score, req))
    scored.sort(key=lambda x: x[0], reverse=True)
    return scored[:top]


WS_CODE_RE = re.compile(r"WS-CE-\d+", re.IGNORECASE)


def _workshop_code(text: str) -> str | None:
    match = WS_CODE_RE.search(text or "")
    return match.group(0).upper() if match else None


def _best_sow_feature(story: dict, features: list[dict]) -> tuple[dict | None, float]:
    blob = f"{story.get('title', '')} {story.get('description', '')}"
    st = _tokens(blob)
    best_feat: dict | None = None
    best_score = 0.0
    for feat in features:
        feat_blob = " ".join(
            str(feat.get(key, ""))
            for key in ("feature_name", "bpm_topic", "workshop_code")
        )
        score = _jaccard(st, _tokens(feat_blob))
        if score > best_score:
            best_score = score
            best_feat = feat
    return best_feat, best_score


def _objective_for_workshop(unified_kb: dict, code: str | None) -> dict | None:
    if not code:
        return None
    for req in unified_kb.get("requirements") or []:
        statement = req.get("statement") or ""
        locator = req.get("source_locator") or ""
        if code.upper() in statement.upper() or code.upper() in locator.upper():
            if req.get("source_type") == "Pre-read Materials" or "OBJ-" in str(req.get("id", "")):
                return req
    return None


def _why_mapped_to_feature(
    story: dict,
    display_feature: str,
    sow_feat: dict | None,
    matches: list[tuple[float, dict]],
    refs: list[dict],
    unified_kb: dict,
) -> str:
    code = _workshop_code(display_feature) or (
        sow_feat.get("workshop_code") if sow_feat else None
    )
    objective = _objective_for_workshop(unified_kb, code)
    bpm = (sow_feat or {}).get("bpm_topic") or display_feature
    workshop = (sow_feat or {}).get("workshop_code") or code or ""

    evidence = None
    for _score, req in matches:
        if workshop and workshop.upper() in (
            (req.get("statement") or "") + (req.get("source_locator") or "")
        ).upper():
            evidence = req
            break
    if evidence is None and refs and has_supporting_refs(refs):
        evidence = {
            "source_type": refs[0].get("source_type"),
            "source_locator": refs[0].get("locator"),
            "statement": refs[0].get("excerpt"),
        }
    if objective and not evidence:
        evidence = objective

    title = (story.get("title") or "").strip()
    if evidence:
        stype = evidence.get("source_type") or "Feature List"
        loc = evidence.get("source_locator") or evidence.get("locator") or ""
        excerpt = (evidence.get("statement") or evidence.get("excerpt") or "")[:160]
        sow_ref = f"{workshop} — {bpm}".strip(" —") if workshop or bpm else bpm
        return (
            f"Story \"{title}\" aligns with SOW BPM topic {sow_ref}. "
            f"{stype} ({loc}): {excerpt}"
        )
    if sow_feat and sow_feat.get("bpm_topic"):
        return (
            f"Story \"{title}\" shares scope with SOW annexure {sow_feat.get('workshop_code', '')} "
            f"({sow_feat.get('bpm_topic')}) per {sow_feat.get('source_locator', 'SOW.json')}."
        )
    return (
        f"No workshop source cites a clear SOW BPM link for \"{title}\"; "
        f"feature label \"{display_feature}\" is not source-justified."
    )


def _feature_name_and_rationale(
    story: dict,
    features: list[dict],
    unified_kb: dict,
    matches: list[tuple[float, dict]],
    refs: list[dict],
) -> tuple[str, str]:
    feat_ctx = story.get("feature") or {}
    csv_feature = (feat_ctx.get("title") or "").strip()
    sow_feat, _score = _best_sow_feature(story, features)
    if csv_feature:
        display = csv_feature
    elif sow_feat:
        display = sow_feat.get("feature_name", "") or sow_feat.get("bpm_topic", "Unmapped")
    else:
        display = "Unmapped"
    if sow_feat is None and display != "Unmapped":
        for feat in features:
            if _jaccard(_tokens(display), _tokens(feat.get("feature_name", ""))) > 0.2:
                sow_feat = feat
                break
    why = _why_mapped_to_feature(story, display, sow_feat, matches, refs, unified_kb)
    return display, why


def _requirements_for_feature(unified_kb: dict, feat: dict) -> list[dict]:
    code = (feat.get("workshop_code") or "").upper()
    bpm_tokens = _tokens(feat.get("bpm_topic") or "")
    seen: set[str] = set()
    out: list[dict] = []
    for req in unified_kb.get("requirements") or []:
        stmt = req.get("statement") or ""
        loc = req.get("source_locator") or ""
        req_id = str(req.get("id") or "")
        key = req_id or loc or stmt[:60]
        if key in seen:
            continue
        workshop_hits = WS_CODE_RE.findall(stmt)
        if len(workshop_hits) > 1:
            continue
        if req_id.startswith("OBJ-") and code and code in stmt.upper():
            out.append(req)
            seen.add(key)
            continue
        hay = f"{stmt} {loc}".upper()
        if code and code in hay and len(workshop_hits) <= 1:
            if req.get("source_type") == "Pre-read Materials" and len(stmt) > 280:
                continue
            out.append(req)
            seen.add(key)
        elif bpm_tokens and _jaccard(bpm_tokens, _tokens(stmt)) >= 0.25:
            out.append(req)
            seen.add(key)
    objective = _objective_for_workshop(unified_kb, code)
    if objective:
        okey = objective.get("id") or objective.get("source_locator")
        if okey not in seen:
            out.insert(0, objective)
            seen.add(str(okey))
    return out


def _story_covers_requirement(story_blobs: list[str], req: dict, threshold: float = 0.12) -> bool:
    req_tokens = _tokens(req.get("statement") or "")
    if not req_tokens:
        return True
    for blob in story_blobs:
        if _jaccard(req_tokens, _tokens(blob)) >= threshold:
            return True
    return False


def _feature_code_from_display(name: str) -> str | None:
    return _workshop_code(name)


def _why_generated_detail(
    story: dict,
    matches: list[tuple[float, dict]],
    refs: list[dict],
    match_type: str,
    side: str,
) -> str:
    title = (story.get("title") or "").strip()
    description = (story.get("description") or "").strip()
    ac_lines = story.get("acceptance_criteria") or []
    ac_preview = "; ".join(ac_lines[:3])
    if len(ac_lines) > 3:
        ac_preview += f" … (+{len(ac_lines) - 3} more)"

    lines: list[str] = []
    lines.append(
        f"This user story — \"{title}\" — documents a workshop capability or process step "
        f"traced against the unified knowledge base ({match_type} to indexed sources)."
    )
    if description:
        desc_one = " ".join(description.split())
        if len(desc_one) > 320:
            desc_one = desc_one[:317] + "…"
        lines.append(f"Description scope: {desc_one}")
    if ac_lines:
        lines.append(
            f"Acceptance criteria ({len(ac_lines)}): {ac_preview}"
        )

    if has_supporting_refs(refs):
        sources = ", ".join(sorted({r.get("source_type", "") for r in refs if r.get("source_type")}))
        lines.append(f"Primary evidence types: {sources}.")
        for ref in refs[:4]:
            excerpt = (ref.get("excerpt") or "").replace("\n", " ")
            if len(excerpt) > 200:
                excerpt = excerpt[:197] + "…"
            lines.append(
                f"• {ref.get('source_type', '')} ({ref.get('locator', '')}): {excerpt}"
            )
    else:
        lines.append(
            "No supporting workshop excerpt was found in the unified knowledge base for this story text."
        )

    if side == "generated":
        lines.append(
            "This row comes from the AI-generated transcript intake workbook. "
            "The citations above are why the story belongs in that deliverable relative to workshop sources."
        )
    else:
        lines.append(
            "This row comes from the manual consultant workbook. "
            "Generation rationale is retained in JSON but omitted from the manual HTML report per operator policy."
        )

    return "\n\n".join(lines)


def _build_feature_traceability_matrix(
    stories_out: list[dict],
    packet_stories: list[dict],
    features: list[dict],
    unified_kb: dict,
) -> tuple[list[dict], list[dict]]:
    blobs_by_code: dict[str, list[str]] = {}
    titles_by_code: dict[str, list[str]] = {}
    for s_out, s_in in zip(stories_out, packet_stories, strict=False):
        code = _feature_code_from_display(s_out.get("feature_name", ""))
        if not code:
            sow_feat, _ = _best_sow_feature(s_in, features)
            code = (sow_feat or {}).get("workshop_code") or "UNMAPPED"
        code = code.upper() if code != "UNMAPPED" else code
        blob = f"{s_in.get('title', '')} {s_in.get('description', '')}"
        blobs_by_code.setdefault(code, []).append(blob)
        titles_by_code.setdefault(code, []).append(s_out.get("user_story_title", ""))

    matrix: list[dict] = []
    flat_gaps: list[dict] = []
    total_stories = max(len(stories_out), 1)

    for feat in features:
        code = (feat.get("workshop_code") or "").upper()
        feat_name = feat.get("feature_name", "")
        mapped_titles = titles_by_code.get(code, [])
        blobs = blobs_by_code.get(code, [])
        reqs = _requirements_for_feature(unified_kb, feat)
        not_converted: list[dict] = []

        for req in reqs:
            if _story_covers_requirement(blobs, req):
                continue
            not_converted.append(
                {
                    "requirement": (req.get("statement") or "")[:400],
                    "source_type": req.get("source_type", "Feature List"),
                    "source_locator": req.get("source_locator", ""),
                    "why_not_converted": (
                        f"Workshop sources record this requirement ({req.get('source_locator', '')}) "
                        f"for {code or feat_name}, but no user story in this CSV overlaps it in "
                        f"title or description—it was not decomposed into a story in this workbook "
                        f"or may belong to another annexure export."
                    ),
                }
            )

        if not blobs and not not_converted:
            objective = _objective_for_workshop(unified_kb, code)
            stmt = (objective or {}).get("statement") or feat.get("bpm_topic") or feat_name
            not_converted.append(
                {
                    "requirement": stmt[:400],
                    "source_type": (objective or {}).get("source_type", "Feature List"),
                    "source_locator": (objective or {}).get("source_locator", feat.get("source_locator", "")),
                    "why_not_converted": (
                        f"No user stories in this CSV are mapped to annexure {code} ({feat_name}). "
                        f"The SOW/pre-read topic was not represented as user stories in this file."
                    ),
                }
            )

        if reqs:
            covered = len(reqs) - len(
                [r for r in reqs if not _story_covers_requirement(blobs, r)]
            )
            coverage_percent = round(100 * covered / len(reqs), 1)
        else:
            coverage_percent = 100.0 if mapped_titles else 0.0

        row = {
            "feature_name": feat_name,
            "workshop_code": code,
            "mapped_user_stories": mapped_titles,
            "coverage_percent": coverage_percent,
            "missing_requirements": [g["requirement"][:120] for g in not_converted],
            "requirements_not_in_stories": not_converted,
        }
        matrix.append(row)
        for gap in not_converted:
            flat_gaps.append(
                {
                    "statement": gap["requirement"],
                    "source_type": gap["source_type"],
                    "source_path": feat.get("source_path", ""),
                    "source_locator": gap["source_locator"],
                    "confidence": "explicit",
                    "why_not_converted": gap["why_not_converted"],
                    "feature_name": feat_name,
                }
            )

    return matrix, flat_gaps


def _analyze_side(
    packet: dict,
    unified_kb: dict,
    side: str,
    csv_path: str,
    kb_path: str,
) -> dict:
    features = unified_kb.get("features") or []
    stories_out: list[dict] = []
    exact = partial = no_match = 0
    fully = partial_cov = unsupported = 0

    for story in packet.get("user_stories") or []:
        blob = f"{story.get('title', '')}\n{story.get('description', '')}\n" + "\n".join(
            story.get("acceptance_criteria") or []
        )
        matches = _best_matches(blob, unified_kb)

        derived: list[str] = []
        refs: list[dict] = []
        for score, req in matches:
            if score <= 0:
                continue
            stype = req.get("source_type") or "Transcript"
            if stype not in derived:
                derived.append(stype)
            refs.append(
                {
                    "source_type": stype,
                    "locator": req.get("source_locator", ""),
                    "excerpt": (req.get("statement") or "")[:280],
                }
            )
        if not has_supporting_refs(refs):
            refs = [
                {
                    "source_type": "Transcript",
                    "locator": "n/a",
                    "excerpt": NO_EVIDENCE_EXCERPT,
                }
            ]
            derived = []

        best_score = matches[0][0] if matches else 0.0
        match_type = match_type_from_evidence(has_supporting_refs(refs), best_score)
        feature_name, why_feat = _feature_name_and_rationale(
            story, features, unified_kb, matches, refs
        )

        if match_type == "Exact Match":
            exact += 1
        elif match_type == "Partial Match":
            partial += 1
        else:
            no_match += 1

        if match_type == "Exact Match":
            cov = "Fully Supported"
            fully += 1
        elif match_type == "Partial Match":
            cov = "Partially Supported"
            partial_cov += 1
        else:
            cov = "Unsupported"
            unsupported += 1

        conf = f"{matches[0][0]:.2f}" if matches else "0.00"
        action = (
            "Accept"
            if cov == "Fully Supported"
            else ("Clarify with workshop" if cov == "Partially Supported" else "Remove or rewrite")
        )
        stories_out.append(
            {
                "story_key": story.get("story_key"),
                "user_story_id": story.get("id", ""),
                "user_story_title": story.get("title", ""),
                "feature_name": feature_name,
                "match_type": match_type,
                "derived_sources": derived,
                "source_references": refs,
                "why_generated": _why_generated_detail(
                    story, matches, refs, match_type, side
                ),
                "why_mapped_to_feature": why_feat,
                "supporting_business_context": (
                    matches[0][1].get("statement", "")[:200] if matches else ""
                ),
                "coverage_assessment": cov,
                "confidence_score": conf,
                "recommended_action": action,
            }
        )

    total = len(stories_out)
    trace = round(100 * (fully + 0.5 * partial_cov) / total, 1) if total else 0

    matrix, gap_reqs = _build_feature_traceability_matrix(
        stories_out,
        packet.get("user_stories") or [],
        features,
        unified_kb,
    )

    return {
        "schema_version": "1.0",
        "kind": "user-story-source-analysis",
        "side": side,
        "status": "Draft — Pending Method Review",
        "assumption_id": "ASSUMPTION-USHTML-001",
        "inputs": {"user_story_csv": csv_path, "unified_knowledge_base": kb_path},
        "executive_summary": {
            "total_user_stories": total,
            "exact_matches": exact,
            "partial_matches": partial,
            "no_matches": no_match,
            "fully_supported_stories": fully,
            "overall_traceability_score": str(trace),
        },
        "stories": stories_out,
        "feature_traceability_matrix": matrix,
        "source_coverage_summary": {
            "transcript_percent": round(
                100
                * sum(1 for s in stories_out if "Transcript" in (s.get("derived_sources") or []))
                / max(total, 1),
                1,
            ),
            "mom_percent": round(
                100
                * sum(1 for s in stories_out if "MoM" in (s.get("derived_sources") or []))
                / max(total, 1),
                1,
            ),
            "preread_percent": round(
                100
                * sum(
                    1 for s in stories_out if "Pre-read Materials" in (s.get("derived_sources") or [])
                )
                / max(total, 1),
                1,
            ),
            "feature_list_percent": round(
                100
                * sum(1 for s in stories_out if s.get("feature_name") != "Unmapped")
                / max(total, 1),
                1,
            ),
        },
        "coverage_gap_analysis": {
            "feature_requirements_not_in_stories": gap_reqs,
            "source_requirements_not_captured": [],
            "unmapped_user_stories": [
                {
                    "statement": s["user_story_title"],
                    "source_type": "User Story CSV",
                    "source_path": csv_path,
                    "source_locator": s.get("story_key", ""),
                    "confidence": "explicit",
                }
                for s in stories_out
                if s.get("feature_name") == "Unmapped" or s.get("match_type") == "No Match"
            ],
            "missing_business_scenarios": [],
        },
        "open_items": [
            "ASSUMPTION-USHTML-001 pending method review.",
            "MoM PDF intake-safety exit 1 (card-number-shaped false positives); run used normalized mom .txt.",
        ],
    }


def main() -> int:
    alloc_file = PIPE / "2026-10-16-marketing-ws03-allocated-paths.json"
    if not alloc_file.is_file():
        print("FAIL: missing allocated paths JSON", file=sys.stderr)
        return 1
    paths = json.loads(alloc_file.read_text(encoding="utf-8"))

    manual_csv = PIPE / "2026-10-16-marketing-ws03-manual.csv"
    ai_csv = PIPE / "2026-10-16-marketing-ws03-generated.csv"
    _write_csv(manual_csv, _rulebook_rows_from_xlsx(ROOT / "CRM - Workshop Stories.xlsx", legacy_manual=True))
    _write_csv(
        ai_csv,
        _rulebook_rows_from_xlsx(
            ROOT / "artifacts/intake/2026-10-13-marketing-ws03-user-stories-from-transcript.xlsx"
        ),
    )

    transcript = ROOT / "artifacts/intake/2026-10-13-marketing-ws03-transcript.txt"
    mom = INTAKE_PIPE / "2026-10-16-marketing-ws03-mom.txt"
    preread = INTAKE_PIPE / "2026-10-13-marketing-ws03-prerequisite.txt"
    sow = ROOT / "artifacts/work-items/SOW.json"

    _run(
        [
            sys.executable,
            str(SCRIPTS / "generate_user_story_source_analysis.py"),
            "prepare-kb",
            "--transcript",
            str(transcript),
            "--mom",
            str(mom),
            "--preread",
            str(preread),
            "--sow-json",
            str(sow),
            "--out",
            paths["kb_index"],
        ]
    )
    kb_index = json.loads(Path(paths["kb_index"]).read_text(encoding="utf-8"))
    unified = _build_unified_kb(
        kb_index, ROOT / "artifacts/intake/2026-10-13-marketing-ws03-workshop-prerequisite-knowledge-base.json"
    )
    Path(paths["unified_kb"]).write_text(json.dumps(unified, indent=2), encoding="utf-8")

    for side, csv, packet_key, analysis_key in (
        ("manual", manual_csv, "manual_packet", "manual_analysis"),
        ("generated", ai_csv, "ai_packet", "ai_analysis"),
    ):
        _run(
            [
                sys.executable,
                str(SCRIPTS / "generate_user_story_source_analysis.py"),
                "prepare-stories",
                "--csv",
                str(csv),
                "--side",
                side,
                "--kb-index",
                paths["kb_index"],
                "--unified-kb",
                paths["unified_kb"],
                "--out",
                paths[packet_key],
            ]
        )
        packet = json.loads(Path(paths[packet_key]).read_text(encoding="utf-8"))
        analysis = _analyze_side(packet, unified, side, str(csv), paths["unified_kb"])
        Path(paths[analysis_key]).write_text(json.dumps(analysis, indent=2), encoding="utf-8")
        pub_html = paths["public_manual_html"] if side == "manual" else paths["public_ai_html"]
        pub_json = paths["public_manual_json"] if side == "manual" else paths["public_ai_json"]
        _run(
            [
                sys.executable,
                str(SCRIPTS / "generate_user_story_source_analysis.py"),
                "validate",
                "--analysis",
                paths[analysis_key],
            ]
        )
        _run(
            [
                sys.executable,
                str(SCRIPTS / "generate_user_story_source_analysis.py"),
                "render",
                "--analysis",
                paths[analysis_key],
                "--html",
                pub_html,
                "--json",
                pub_json,
            ]
        )

    alloc_file.write_text(json.dumps(paths, indent=2), encoding="utf-8")
    print("DONE", paths["public_manual_html"], paths["public_ai_html"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
