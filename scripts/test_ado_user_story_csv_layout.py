#!/usr/bin/env python3
"""Regression tests for ADO CSV/xlsx layout (run: python scripts/test_ado_user_story_csv_layout.py)."""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import yaml

from ado_user_story_csv_layout import (
    acceptance_criteria_cell,
    apply_generation_defaults,
    headers_from_rulebook,
    stories_sheet_name,
    validate_structure,
    write_csv,
    write_xlsx,
)
from importlib.machinery import SourceFileLoader

_validate = SourceFileLoader(
    "validate_user_story_csv",
    str(ROOT / "scripts" / "validate-user-story-csv.py"),
).load_module()


def _rulebook() -> dict:
    return yaml.safe_load((ROOT / "docs/intake/user-story-csv-rulebook.yaml").read_text(encoding="utf-8"))


def _headers(rulebook: dict) -> list[str]:
    return headers_from_rulebook(rulebook)


def _row(rulebook: dict, wit: str, **fields: str) -> dict[str, str]:
    headers = _headers(rulebook)
    base = {h: "" for h in headers}
    base["Work Item Type"] = wit
    if "Title" in fields:
        base["Title"] = fields.pop("Title")
    base.update(fields)
    apply_generation_defaults(rulebook, base)
    if "Title" not in headers:
        base.pop("Title", None)
    return base


def test_headers_match_excel_title_split() -> None:
    rulebook = _rulebook()
    headers = _headers(rulebook)
    assert headers == rulebook["csv"]["column_order"]
    assert headers == rulebook["columns"]["stories"]["order"]
    assert "Title" not in headers
    assert headers[:6] == ["State", "ID", "Work Item Type", "Title 1", "Title 2", "Title 3"]
    assert "Assigned To" in headers
    assert "User Story Type" in headers
    assert "Task Type1" not in headers
    assert "Start Date" not in headers


def test_good_export_passes() -> None:
    rulebook = _rulebook()
    headers = _headers(rulebook)
    rows = [
        _row(
            rulebook,
            "Workshop",
            **{
                "Title 1": "WS-01 - Sample workshop",
                "Description": "Session scope",
                "Area Path": r"Day 2\Area",
            },
        ),
        _row(
            rulebook,
            "User Story",
            **{
                "Title 3": "Capture the record",
                "Description": "As a user, I need to capture the record, so that downstream work can start.",
                "Acceptance Criteria": acceptance_criteria_cell("The record is saved", "A duplicate is rejected"),
            },
        ),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        csv_path = Path(tmp) / "working.csv"
        write_csv(csv_path, headers, rows)
        file_headers, file_rows = _validate.read_csv_rows(csv_path)
        result = _validate.validate(rulebook, file_headers, file_rows)
        assert result["status"] == "PASS", result


def test_user_story_area_path_allowed() -> None:
    rulebook = _rulebook()
    rows = [
        _row(
            rulebook,
            "Workshop",
            **{"Title 1": "WS-01 - Test", "Description": "Scope", "Area Path": r"Day 2\Area"},
        ),
        _row(
            rulebook,
            "User Story",
            **{
                "Title 3": "Capture the record",
                "Description": "As a user, I need to capture the record, so that value is delivered.",
                "Area Path": r"Day 2\Area",
                "Acceptance Criteria": acceptance_criteria_cell("One", "Two"),
            },
        ),
    ]
    findings, critical = validate_structure(rulebook, rows)
    assert not any("Area Path must be empty" in f for f in findings), findings
    assert critical == 0, findings


def test_legacy_title_projects_to_title_3() -> None:
    rulebook = _rulebook()
    headers = _headers(rulebook)
    row = {h: "" for h in headers}
    row["Work Item Type"] = "User Story"
    row["Title"] = "Capture the record"
    apply_generation_defaults(rulebook, row)
    assert row["Title 3"] == "Capture the record"
    assert row.get("Title 1", "") == ""
    assert row.get("Title 2", "") == ""


def test_single_line_ac_fails() -> None:
    rulebook = _rulebook()
    bad = [
        _row(
            rulebook,
            "Workshop",
            **{"Title 1": "WS-01 - Test Workshop", "Description": "Scope", "Area Path": r"Day 2\Area"},
        ),
        _row(
            rulebook,
            "User Story",
            **{
                "Title 3": "Bad AC",
                "Description": "As a user, I need to test AC layout, so that ADO import works.",
                "Acceptance Criteria": "- One - Two on one line",
            },
        ),
    ]
    findings, critical = validate_structure(rulebook, bad)
    assert critical > 0
    assert any("one '- ' bullet per line" in f for f in findings) or any(
        "multiple on one line" in f for f in findings
    )


def test_traceable_test_case_passes() -> None:
    rulebook = _rulebook()
    rows = [
        _row(rulebook, "Epic", **{"ID": "67476", "Title 1": "16. Catalog epic", "Description": "Catalog epic"}),
        _row(rulebook, "Feature", **{"ID": "67535", "Title 2": "Catalog feature", "Description": "Catalog feature"}),
        _row(
            rulebook,
            "Test Case",
            **{
                "Title 3": "Order waits for approval",
                "Description": (
                    "Transcript: artifacts/intake/session.txt#00:12:00\n"
                    "Epic: 67476 | 16. Catalog epic\n"
                    "Feature: 67535 | Catalog feature\n"
                    "Steps:\n"
                    "1. Submit a sales order that the transcript says must be approved"
                ),
                "Acceptance Criteria": acceptance_criteria_cell("The order remains pending approval"),
            },
        ),
    ]
    findings, critical = validate_structure(rulebook, rows)
    assert critical == 0, findings
    result = _validate.validate(rulebook, _headers(rulebook), rows)
    assert result["status"] == "PASS", result


def test_public_xlsx_uses_rulebook_styling() -> None:
    from openpyxl import load_workbook

    rulebook = _rulebook()
    headers = _headers(rulebook)
    rows = [
        _row(
            rulebook,
            "Workshop",
            **{"Title 1": "WS-01 - Sample", "Description": "Scope", "Area Path": r"Day 2\Area"},
        ),
        _row(
            rulebook,
            "User Story",
            **{
                "Title 3": "Capture the record",
                "Description": "As a user, I need to capture the record, so that work can start.",
                "Acceptance Criteria": acceptance_criteria_cell("Saved"),
            },
        ),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        write_xlsx(xlsx, headers, rows, rulebook=rulebook, project_label="")
        wb = load_workbook(xlsx)
        assert wb.active.title == stories_sheet_name(rulebook, "")
        ws = wb.active
        assert [ws.cell(1, c).value for c in range(1, len(headers) + 1)] == headers
        assert ws.cell(2, 3).value == "Workshop"
        title1_col = headers.index("Title 1") + 1
        assert ws.cell(2, title1_col).value == "WS-01 - Sample"
        title3_col = headers.index("Title 3") + 1
        assert ws.cell(3, title3_col).value == "Capture the record"
        assigned_col = headers.index("Assigned To") + 1
        from openpyxl.utils import get_column_letter

        assert ws.column_dimensions[get_column_letter(assigned_col)].hidden is True
        fill = ws.cell(2, 3).fill
        rgb = getattr(fill.fgColor, "rgb", None)
        assert rgb in {"FFFFFF00", "FFFF00"} or str(rgb).endswith("FFFF00")


def main() -> int:
    test_headers_match_excel_title_split()
    test_good_export_passes()
    test_user_story_area_path_allowed()
    test_legacy_title_projects_to_title_3()
    test_single_line_ac_fails()
    test_traceable_test_case_passes()
    test_public_xlsx_uses_rulebook_styling()
    print("ado_user_story_csv_layout tests OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
