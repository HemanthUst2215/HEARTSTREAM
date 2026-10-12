#!/usr/bin/env python3
"""Canonical MSD E2E test-case workbook layout: role lookup, write helpers, validation.

All `/generate-test-cases-from-user-stories` Excel output must be written through
`write_workbook()` and pass `validate_workbook()`. Column names, sheet names,
ID patterns, and enums come from docs/intake/msd-test-case-rulebook.yaml.
"""
from __future__ import annotations

import csv
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

try:
    import yaml
except ImportError:  # pragma: no cover
    yaml = None  # type: ignore

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore
    load_workbook = None  # type: ignore
    Alignment = None  # type: ignore
    Font = None  # type: ignore
    PatternFill = None  # type: ignore
    get_column_letter = None  # type: ignore
    Worksheet = Any  # type: ignore


def load_rulebook(path: Path) -> dict[str, Any]:
    if yaml is None:
        raise SystemExit("FAIL: PyYAML required (pip install pyyaml)")
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise SystemExit("rulebook root must be a mapping")
    return data


def role(rulebook: dict[str, Any], sheet_key: str, role_name: str) -> str:
    roles = rulebook["columns"][sheet_key]["roles"]
    header = roles.get(role_name)
    if not header:
        raise KeyError(f"rulebook columns.{sheet_key}.roles.{role_name} missing")
    return str(header)


def headers(rulebook: dict[str, Any], sheet_key: str) -> list[str]:
    cols = rulebook["columns"][sheet_key].get("order")
    if not cols:
        raise ValueError(f"rulebook columns.{sheet_key}.order missing")
    return list(cols)


def sheet_cfg(rulebook: dict[str, Any], sheet_key: str) -> dict[str, Any]:
    return rulebook["workbook"]["sheets"][sheet_key]


def sheet_name(rulebook: dict[str, Any], sheet_key: str) -> str:
    return str(sheet_cfg(rulebook, sheet_key)["name"])


def project_label(rulebook: dict[str, Any], config_name: str | None = None) -> str:
    spec = rulebook.get("workbook", {}).get("project_label", {})
    placeholders = set(spec.get("placeholder_tokens") or ["<project-name>", ""])
    name = (config_name or "").strip()
    if name and name not in placeholders:
        return name
    return str(spec.get("fallback") or "E2E")


def format_test_case_id(rulebook: dict[str, Any], scenario_id: int, polarity: str, seq: int = 1) -> str:
    spec = rulebook["id_patterns"]["test_case_id"]
    polar = spec["polarity"][polarity]
    return str(spec["pattern"]).format(scenario=int(scenario_id), polarity=polar, seq=int(seq))


def test_case_id_re(rulebook: dict[str, Any]) -> re.Pattern[str]:
    raw = rulebook["id_patterns"]["test_case_id"].get("regex") or r"^TC(\d{2})-([PN])(\d{2})$"
    return re.compile(raw)


def blank_row(column_order: list[str]) -> dict[str, str]:
    return {h: "" for h in column_order}


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def join_unique(values: list[str], sep: str) -> str:
    seen: list[str] = []
    for raw in values:
        for part in (raw or "").split(sep):
            token = part.strip()
            if token and token not in seen:
                seen.append(token)
    return sep.join(seen)


def _require_openpyxl() -> None:
    if Workbook is None or load_workbook is None:
        raise SystemExit("FAIL: openpyxl required (pip install openpyxl)")


def _styling(rulebook: dict[str, Any]) -> dict[str, Any]:
    return dict(rulebook.get("workbook", {}).get("styling") or {})


def _style_colors(styling: dict[str, Any]) -> dict[str, str]:
    defaults = {
        "title_band": "16324F",
        "table_header": "2F75B5",
        "metrics_header": "2F75B5",
        "metrics_row": "EAF2F8",
        "design_basis_row": "F3F6F8",
        "on_dark": "FFFFFF",
    }
    merged = {**defaults, **(styling.get("colors") or {})}
    return merged


def _argb_hex(color: str) -> str:
    token = str(color).lstrip("#").upper()
    if len(token) == 8:
        return token
    return f"FF{token}"


def _solid_fill(color: str) -> PatternFill:
    return PatternFill(patternType="solid", fgColor=_argb_hex(color))


def _make_font(
    styling: dict[str, Any],
    *,
    size: float | None = None,
    bold: bool = False,
    color: str | None = None,
) -> Font:
    family = str(styling.get("font_family") or "Aptos Narrow")
    body = float(styling.get("body_font_size") or 12)
    kwargs: dict[str, Any] = {"name": family, "bold": bold, "size": size if size is not None else body}
    if color:
        kwargs["color"] = _argb_hex(color)
    return Font(**kwargs)


def _sheet_style(styling: dict[str, Any], sheet_key: str) -> dict[str, Any]:
    sheets = styling.get("sheets") or {}
    return dict(sheets.get(sheet_key) or {})


def _column_widths_for_sheet(styling: dict[str, Any], sheet_key: str, cols: list[str]) -> list[float | None]:
    sheet = _sheet_style(styling, sheet_key)
    by_header = sheet.get("column_widths_by_header") or {}
    default_w = sheet.get("default_column_width")
    out: list[float | None] = []
    for h in cols:
        if h in by_header:
            out.append(float(by_header[h]))
        elif default_w is not None:
            out.append(float(default_w))
        else:
            out.append(None)
    return out


def _merge_title(
    ws: Worksheet,
    row: int,
    col_count: int,
    text: str,
    styling: dict[str, Any],
    *,
    font_size: float,
    bold: bool = True,
    subtitle: bool = False,
) -> None:
    colors = _style_colors(styling)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row, 1, text)
    cell.font = _make_font(styling, size=font_size, bold=bold, color=colors["on_dark"])
    cell.fill = _solid_fill(colors["title_band"])
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    if not subtitle and styling.get("title_row_height"):
        ws.row_dimensions[row].height = float(styling["title_row_height"])


def _write_table_header_row(
    ws: Worksheet,
    row: int,
    cols: list[str],
    styling: dict[str, Any],
    sheet_key: str,
) -> None:
    colors = _style_colors(styling)
    sheet = _sheet_style(styling, sheet_key)
    header_spec = dict(sheet.get("table_header") or {})
    mode = str(header_spec.get("mode") or "branded")
    body = float(styling.get("body_font_size") or 12)
    for i, name in enumerate(cols, start=1):
        cell = ws.cell(row, i, name)
        if mode == "plain":
            cell.font = _make_font(styling, size=body, bold=bool(header_spec.get("bold", False)))
            cell.alignment = Alignment(
                wrap_text=bool(header_spec.get("wrap_text", False)),
                vertical=str(header_spec.get("vertical") or "top"),
            )
        else:
            cell.font = _make_font(
                styling,
                size=float(header_spec.get("font_size") or body),
                bold=True,
                color=colors["on_dark"],
            )
            cell.fill = _solid_fill(colors["table_header"])
            cell.alignment = Alignment(
                wrap_text=bool(header_spec.get("wrap_text", False)),
                vertical=str(header_spec.get("vertical") or "center"),
            )


def _apply_data_row_style(
    ws: Worksheet,
    row: int,
    col_count: int,
    styling: dict[str, Any],
    wrap_col_indexes: list[int],
) -> None:
    body = float(styling.get("body_font_size") or 12)
    wrap_set = set(wrap_col_indexes)
    for c in range(1, col_count + 1):
        cell = ws.cell(row, c)
        cell.font = _make_font(styling, size=body, bold=False)
        cell.alignment = Alignment(
            wrap_text=c in wrap_set,
            vertical="top",
        )


def _paint_cells(ws: Worksheet, row: int, col_start: int, col_end: int, fill: PatternFill) -> None:
    for c in range(col_start, col_end + 1):
        ws.cell(row, c).fill = fill


def _style_metrics_block(
    ws: Worksheet,
    styling: dict[str, Any],
    header_row: int,
    data_start: int,
    row_count: int,
    col_start: int,
    header_cols: int,
    data_cols: int,
    *,
    design_basis_span: int = 0,
) -> None:
    colors = _style_colors(styling)
    header_fill = _solid_fill(colors["metrics_header"])
    data_fill = _solid_fill(colors["metrics_row"])
    basis_fill = _solid_fill(colors["design_basis_row"])
    body = float(styling.get("body_font_size") or 12)
    header_span = design_basis_span if design_basis_span else header_cols
    _paint_cells(ws, header_row, col_start, col_start + header_span - 1, header_fill)
    for c in range(col_start, col_start + header_span):
        cell = ws.cell(header_row, c)
        cell.font = _make_font(styling, size=body, bold=True, color=colors["on_dark"])
    for r_i in range(row_count):
        r = data_start + r_i
        row_fill = basis_fill if design_basis_span else data_fill
        _paint_cells(ws, r, col_start, col_start + (design_basis_span or data_cols) - 1, row_fill)
        for c in range(col_start, col_start + data_cols):
            cell = ws.cell(r, c)
            cell.font = _make_font(styling, size=body, bold=False)
            cell.alignment = Alignment(wrap_text=bool(design_basis_span and c > col_start), vertical="top")


def denormalize_register_rows(rulebook: dict[str, Any], cases: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Expand case objects (role keys) into one register row dict per step (header keys)."""
    cols = headers(rulebook, "register")
    roles = rulebook["columns"]["register"]["roles"]
    status_default = rulebook["status_field"]["generation_default"]
    exec_at_gen = rulebook.get("execution_columns", {}).get("at_generation", {})
    out: list[dict[str, str]] = []
    for case in cases:
        steps = case.get("steps") or []
        for step in steps:
            row = blank_row(cols)
            for role_name, header in roles.items():
                if role_name in ("actual_result", "defect_id", "comments"):
                    continue
                if role_name in step:
                    row[header] = cell_str(step.get(role_name))
                elif role_name in case:
                    row[header] = cell_str(case.get(role_name))
            if exec_at_gen.get("actual_result") == "empty":
                row[roles["actual_result"]] = ""
            if exec_at_gen.get("defect_id") == "empty":
                row[roles["defect_id"]] = ""
            if exec_at_gen.get("comments") == "empty":
                row[roles["comments"]] = ""
            status_mode = exec_at_gen.get("status", "generation_default")
            if status_mode == "generation_default":
                row[roles["status"]] = status_default
            elif status_mode != "empty":
                row[roles["status"]] = cell_str(case.get("status") or status_default)
            out.append(row)
    return out


def coverage_rows_from_cases(rulebook: dict[str, Any], cases: list[dict[str, Any]]) -> list[dict[str, str]]:
    cols = headers(rulebook, "coverage")
    roles = rulebook["columns"]["coverage"]["roles"]
    reg_roles = rulebook["columns"]["register"]["roles"]
    sep = rulebook.get("application_isv", {}).get("composite_separator") or " / "
    critical_token = "Critical"
    register = denormalize_register_rows(rulebook, cases)

    by_scenario: dict[str, list[dict[str, str]]] = defaultdict(list)
    case_by_id = {cell_str(c.get("test_case_id")): c for c in cases}
    for row in register:
        by_scenario[cell_str(row[reg_roles["scenario_id"]])].append(row)

    polarity_re = test_case_id_re(rulebook)
    out: list[dict[str, str]] = []
    for scen in sorted(by_scenario, key=lambda x: int(x) if str(x).isdigit() else str(x)):
        rows = by_scenario[scen]
        first = rows[0]
        ids = []
        for r in rows:
            cid = cell_str(r[reg_roles["test_case_id"]])
            if cid not in ids:
                ids.append(cid)
        positive = ""
        exception = ""
        for cid in ids:
            m = polarity_re.match(cid)
            if not m:
                continue
            if m.group(2) == rulebook["id_patterns"]["test_case_id"]["polarity"]["positive"]:
                if not positive:
                    positive = cid
            else:
                if not exception:
                    exception = cid
        apps = join_unique([r[reg_roles["application_isv"]] for r in rows], sep)
        crit = sum(1 for r in rows if cell_str(r[reg_roles["priority"]]) == critical_token)
        sample = case_by_id.get(positive) or case_by_id.get(exception) or {}
        row = blank_row(cols)
        row[roles["scenario_id"]] = scen
        row[roles["business_process"]] = cell_str(first[reg_roles["business_process"]])
        row[roles["positive_case_id"]] = positive
        row[roles["exception_case_id"]] = exception
        row[roles["detailed_step_count"]] = str(len(rows))
        row[roles["critical_step_count"]] = str(crit)
        row[roles["primary_systems"]] = apps
        row[roles["source_flow"]] = cell_str(sample.get("source_flow"))
        row[roles["coverage_focus"]] = cell_str(sample.get("coverage_focus"))
        out.append(row)
    return out


def isv_rows_from_register(rulebook: dict[str, Any], register_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    cols = headers(rulebook, "isv_lookup")
    roles = rulebook["columns"]["isv_lookup"]["roles"]
    reg = rulebook["columns"]["register"]["roles"]
    sep = rulebook.get("application_isv", {}).get("composite_separator") or " / "
    owners: dict[str, Counter[str]] = defaultdict(Counter)
    for row in register_rows:
        app = cell_str(row.get(reg["application_isv"]))
        owner = cell_str(row.get(reg["step_owner"]))
        for token in [p.strip() for p in app.split(sep) if p.strip()]:
            if owner:
                owners[token][owner] += 1
            else:
                owners[token]  # ensure key
    out: list[dict[str, str]] = []
    for token in sorted(owners):
        row = blank_row(cols)
        row[roles["application"]] = token
        majority = owners[token].most_common(1)
        row[roles["typical_owner"]] = majority[0][0] if majority else ""
        out.append(row)
    return out


def metric_value(formula: str, register_rows: list[dict[str, str]], rulebook: dict[str, Any], status_value: str = "") -> int:
    roles = rulebook["columns"]["register"]["roles"]
    if formula == "distinct_scenario_id":
        return len({cell_str(r[roles["scenario_id"]]) for r in register_rows if cell_str(r[roles["scenario_id"]])})
    if formula == "distinct_test_case_id":
        return len({cell_str(r[roles["test_case_id"]]) for r in register_rows if cell_str(r[roles["test_case_id"]])})
    if formula == "count_register_rows":
        return len(register_rows)
    if formula == "count_status":
        return sum(1 for r in register_rows if cell_str(r[roles["status"]]) == status_value)
    raise ValueError(f"unknown metric formula: {formula}")


def _resolve_text_ref(rulebook: dict[str, Any], ref: str) -> str:
    cur: Any = rulebook
    for part in ref.split("."):
        cur = cur[part]
    return str(cur)


def write_workbook(
    path: Path,
    rulebook: dict[str, Any],
    cases: list[dict[str, Any]],
    *,
    project: str | None = None,
    also_csv: Path | None = None,
) -> dict[str, Any]:
    _require_openpyxl()
    styling = _styling(rulebook)
    label = project_label(rulebook, project)
    register_rows = denormalize_register_rows(rulebook, cases)
    cov_rows = coverage_rows_from_cases(rulebook, cases)
    isv_rows = isv_rows_from_register(rulebook, register_rows)
    scenario_count = metric_value("distinct_scenario_id", register_rows, rulebook)

    wb = Workbook()
    _write_coverage_sheet(wb, rulebook, styling, cov_rows, register_rows, label, scenario_count)
    _write_register_sheet(wb, rulebook, styling, register_rows, label, scenario_count)
    _write_isv_sheet(wb, rulebook, styling, isv_rows)

    default = wb.sheetnames[0]
    if default not in {sheet_name(rulebook, k) for k in ("coverage", "register", "isv_lookup")}:
        if default in wb.sheetnames and wb[default].max_row == 1 and wb[default].max_column == 1:
            del wb[default]

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

    if also_csv is not None:
        write_register_csv(also_csv, headers(rulebook, "register"), register_rows)

    return {
        "xlsx": str(path),
        "register_rows": len(register_rows),
        "cases": len(cases),
        "scenarios": scenario_count,
        "isv_rows": len(isv_rows),
    }


def _write_coverage_sheet(
    wb: Workbook,
    rulebook: dict[str, Any],
    styling: dict[str, Any],
    cov_rows: list[dict[str, str]],
    register_rows: list[dict[str, str]],
    label: str,
    scenario_count: int,
) -> None:
    cfg = sheet_cfg(rulebook, "coverage")
    cols = headers(rulebook, "coverage")
    ws = wb.active
    ws.title = str(cfg["name"])
    n = len(cols)
    cov_sheet = _sheet_style(styling, "coverage")
    title_size = float(cov_sheet.get("title_font_size") or styling.get("title_font_size") or 18)
    subtitle_size = float(cov_sheet.get("subtitle_font_size") or 12)
    _merge_title(
        ws,
        int(cfg["title_row"]),
        n,
        str(cfg["title_template"]).format(project_label=label, scenario_count=scenario_count),
        styling,
        font_size=title_size,
        bold=True,
    )
    _merge_title(
        ws,
        int(cfg["subtitle_row"]),
        n,
        str(cfg["subtitle_template"]).format(project_label=label, scenario_count=scenario_count),
        styling,
        font_size=subtitle_size,
        bold=False,
        subtitle=True,
    )
    _write_metrics(ws, rulebook, styling, register_rows)
    header_row = int(cfg["table_header_row"])
    _write_table_header_row(ws, header_row, cols, styling, "coverage")
    wrap_headers = {role(rulebook, "coverage", r) for r in cov_sheet.get("wrap_roles", [])}
    wrap_idx = [i for i, h in enumerate(cols, start=1) if h in wrap_headers]
    start = int(cfg["data_start_row"])
    for i, row in enumerate(cov_rows):
        r = start + i
        for c, h in enumerate(cols, start=1):
            ws.cell(r, c, row.get(h, ""))
        _apply_data_row_style(ws, r, n, styling, wrap_idx)
    if cfg.get("freeze_panes"):
        ws.freeze_panes = cfg["freeze_panes"]
    for i, w in enumerate(_column_widths_for_sheet(styling, "coverage", cols), start=1):
        if w is not None:
            ws.column_dimensions[get_column_letter(i)].width = w


def _write_metrics(
    ws: Worksheet,
    rulebook: dict[str, Any],
    styling: dict[str, Any],
    register_rows: list[dict[str, str]],
) -> None:
    metrics = rulebook["metrics"]
    cfg = sheet_cfg(rulebook, "coverage")
    header_row = int(cfg["metrics_header_row"])
    data_start = int(cfg["metrics_data_start_row"])
    blocks: list[tuple[str, int, int, int]] = [
        ("coverage_block", 1, 2, 0),
        ("execution_block", 4, 2, 0),
        ("design_basis_block", 7, 2, 3),
    ]
    for block_key, col, data_cols, basis_span in blocks:
        block = metrics[block_key]
        headers_b = list(block["headers"])
        for i, h in enumerate(headers_b):
            ws.cell(header_row, col + i, h)
        row_count = len(block["rows"])
        for r_i, item in enumerate(block["rows"]):
            ws.cell(data_start + r_i, col, item["label"])
            formula = item.get("formula")
            if formula:
                ws.cell(
                    data_start + r_i,
                    col + 1,
                    metric_value(formula, register_rows, rulebook, item.get("status_value") or ""),
                )
            elif item.get("text_ref"):
                ws.cell(data_start + r_i, col + 1, _resolve_text_ref(rulebook, item["text_ref"]))
            elif item.get("text"):
                ws.cell(data_start + r_i, col + 1, item["text"])
        _style_metrics_block(
            ws,
            styling,
            header_row,
            data_start,
            row_count,
            col,
            len(headers_b),
            data_cols,
            design_basis_span=basis_span,
        )


def _write_register_sheet(
    wb: Workbook,
    rulebook: dict[str, Any],
    styling: dict[str, Any],
    register_rows: list[dict[str, str]],
    label: str,
    scenario_count: int,
) -> None:
    cfg = sheet_cfg(rulebook, "register")
    cols = headers(rulebook, "register")
    ws = wb.create_sheet(str(cfg["name"]))
    n = len(cols)
    reg_sheet = _sheet_style(styling, "register")
    title_size = float(reg_sheet.get("title_font_size") or styling.get("title_font_size") or 18)
    subtitle_size = float(reg_sheet.get("subtitle_font_size") or 10)
    _merge_title(
        ws,
        int(cfg["title_row"]),
        n,
        str(cfg["title_template"]).format(project_label=label, scenario_count=scenario_count),
        styling,
        font_size=title_size,
        bold=True,
    )
    _merge_title(
        ws,
        int(cfg["subtitle_row"]),
        n,
        str(cfg["subtitle_template"]).format(project_label=label, scenario_count=scenario_count),
        styling,
        font_size=subtitle_size,
        bold=False,
        subtitle=True,
    )
    _write_table_header_row(ws, int(cfg["table_header_row"]), cols, styling, "register")
    wrap_headers = {role(rulebook, "register", r) for r in rulebook["columns"]["register"].get("wrap_roles", [])}
    wrap_idx = [i for i, h in enumerate(cols, start=1) if h in wrap_headers]
    start = int(cfg["data_start_row"])
    for i, row in enumerate(register_rows):
        r = start + i
        for c, h in enumerate(cols, start=1):
            ws.cell(r, c, row.get(h, ""))
        _apply_data_row_style(ws, r, n, styling, wrap_idx)
    if cfg.get("freeze_panes"):
        ws.freeze_panes = cfg["freeze_panes"]
    for i, w in enumerate(_column_widths_for_sheet(styling, "register", cols), start=1):
        if w is not None:
            ws.column_dimensions[get_column_letter(i)].width = w


def _write_isv_sheet(
    wb: Workbook,
    rulebook: dict[str, Any],
    styling: dict[str, Any],
    isv_rows: list[dict[str, str]],
) -> None:
    cfg = sheet_cfg(rulebook, "isv_lookup")
    cols = headers(rulebook, "isv_lookup")
    ws = wb.create_sheet(str(cfg["name"]))
    header_row = int(cfg["table_header_row"])
    for i, name in enumerate(cols, start=1):
        cell = ws.cell(header_row, i, name)
        cell.font = Font(bold=True)
    start = int(cfg["data_start_row"])
    body = float(styling.get("body_font_size") or 12)
    for i, row in enumerate(isv_rows):
        r = start + i
        for c, h in enumerate(cols, start=1):
            ws.cell(r, c, row.get(h, ""))
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).font = _make_font(styling, size=body, bold=False)
    if cfg.get("visible") is False:
        ws.sheet_state = "hidden"


def write_register_csv(path: Path, cols: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=cols, lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_sheet_table(
    ws: Worksheet,
    header_row: int,
    data_start_row: int,
    expected_headers: list[str],
) -> tuple[list[str], list[dict[str, str]]]:
    got = [cell_str(ws.cell(header_row, c).value) for c in range(1, len(expected_headers) + 1)]
    rows: list[dict[str, str]] = []
    r = data_start_row
    while r <= ws.max_row:
        values = [ws.cell(r, c).value for c in range(1, len(expected_headers) + 1)]
        if all(v is None or str(v).strip() == "" for v in values):
            # stop at first fully empty row after data start
            if rows:
                break
            r += 1
            continue
        rows.append({expected_headers[i]: cell_str(values[i]) for i in range(len(expected_headers))})
        r += 1
    return got, rows


def load_register_from_xlsx(path: Path, rulebook: dict[str, Any]) -> tuple[list[str], list[dict[str, str]]]:
    _require_openpyxl()
    wb = load_workbook(path, data_only=True)
    cfg = sheet_cfg(rulebook, "register")
    ws = wb[str(cfg["name"])]
    return read_sheet_table(ws, int(cfg["table_header_row"]), int(cfg["data_start_row"]), headers(rulebook, "register"))


def validate_workbook(rulebook: dict[str, Any], xlsx_path: Path) -> tuple[list[str], int]:
    _require_openpyxl()
    findings: list[str] = []
    critical = 0
    wb = load_workbook(xlsx_path, data_only=True)

    expected_names = [sheet_name(rulebook, k) for k in ("coverage", "register", "isv_lookup")]
    for name in expected_names:
        if name not in wb.sheetnames:
            findings.append(f"missing sheet {name!r}")
            critical += 1
    extra = [n for n in wb.sheetnames if n not in expected_names]
    if extra:
        findings.append(f"unexpected sheets: {extra}")
        critical += 1

    if critical:
        return findings, critical

    cfg_reg = sheet_cfg(rulebook, "register")
    cfg_cov = sheet_cfg(rulebook, "coverage")
    cfg_isv = sheet_cfg(rulebook, "isv_lookup")
    ws_reg = wb[sheet_name(rulebook, "register")]
    ws_cov = wb[sheet_name(rulebook, "coverage")]
    ws_isv = wb[sheet_name(rulebook, "isv_lookup")]

    if cfg_isv.get("visible") is False and ws_isv.sheet_state != "hidden":
        findings.append(f"sheet {cfg_isv['name']!r} must be hidden")
        critical += 1

    reg_headers, reg_rows = read_sheet_table(
        ws_reg, int(cfg_reg["table_header_row"]), int(cfg_reg["data_start_row"]), headers(rulebook, "register")
    )
    if reg_headers != headers(rulebook, "register"):
        findings.append(f"register header mismatch: expected {headers(rulebook, 'register')!r}, got {reg_headers!r}")
        critical += 1

    cov_headers, cov_rows = read_sheet_table(
        ws_cov, int(cfg_cov["table_header_row"]), int(cfg_cov["data_start_row"]), headers(rulebook, "coverage")
    )
    if cov_headers != headers(rulebook, "coverage"):
        findings.append(f"coverage header mismatch: expected {headers(rulebook, 'coverage')!r}, got {cov_headers!r}")
        critical += 1

    isv_headers, _isv_rows = read_sheet_table(
        ws_isv, int(cfg_isv["table_header_row"]), int(cfg_isv["data_start_row"]), headers(rulebook, "isv_lookup")
    )
    if isv_headers != headers(rulebook, "isv_lookup"):
        findings.append(f"isv header mismatch: expected {headers(rulebook, 'isv_lookup')!r}, got {isv_headers!r}")
        critical += 1

    f2, c2 = validate_register_rows(rulebook, reg_rows)
    findings.extend(f2)
    critical += c2

    expected_cov = coverage_rows_from_cases(
        rulebook,
        _cases_from_register(rulebook, reg_rows),
    )
    # Compare computed detailed step counts to coverage sheet (formula integrity).
    cov_roles = rulebook["columns"]["coverage"]["roles"]
    got_by_scen = {cell_str(r[cov_roles["scenario_id"]]): r for r in cov_rows}
    for exp in expected_cov:
        scen = exp[cov_roles["scenario_id"]]
        got = got_by_scen.get(scen)
        if not got:
            findings.append(f"coverage: missing Scenario ID {scen}")
            critical += 1
            continue
        if exp[cov_roles["detailed_step_count"]] != got[cov_roles["detailed_step_count"]]:
            findings.append(
                f"coverage scenario {scen}: Detailed Steps {got[cov_roles['detailed_step_count']]!r} "
                f"!= register count {exp[cov_roles['detailed_step_count']]!r}"
            )
            critical += 1

    return findings, critical


def _cases_from_register(rulebook: dict[str, Any], rows: list[dict[str, str]]) -> list[dict[str, Any]]:
    roles = rulebook["columns"]["register"]["roles"]
    case_roles = rulebook["columns"]["register"]["case_level_roles"]
    step_roles = rulebook["columns"]["register"]["step_level_roles"]
    grouped: dict[str, dict[str, Any]] = {}
    order: list[str] = []
    for row in rows:
        cid = cell_str(row[roles["test_case_id"]])
        if cid not in grouped:
            case = {k: cell_str(row[roles[k]]) for k in case_roles}
            case["steps"] = []
            grouped[cid] = case
            order.append(cid)
        step = {k: cell_str(row[roles[k]]) for k in step_roles}
        grouped[cid]["steps"].append(step)
    return [grouped[c] for c in order]


def validate_register_rows(rulebook: dict[str, Any], rows: list[dict[str, str]], row_offset: int | None = None) -> tuple[list[str], int]:
    findings: list[str] = []
    critical = 0
    roles = rulebook["columns"]["register"]["roles"]
    required = rulebook["required_fields"]["register_step"]["required_roles"]
    allowed_types = set(rulebook["test_types"]["allowed"])
    allowed_prio = set(rulebook["priority"]["allowed"])
    allowed_status = set(rulebook["status_field"]["allowed"])
    status_default = rulebook["status_field"]["generation_default"]
    id_re = test_case_id_re(rulebook)
    min_steps = int(rulebook["generation"]["steps"].get("min_per_case") or 1)
    case_roles = rulebook["columns"]["register"]["case_level_roles"]
    start = row_offset if row_offset is not None else int(sheet_cfg(rulebook, "register")["data_start_row"])

    if not rows:
        findings.append("register has no data rows")
        return findings, 1

    by_case: dict[str, list[tuple[int, dict[str, str]]]] = defaultdict(list)
    for i, row in enumerate(rows):
        line = start + i
        cid = cell_str(row.get(roles["test_case_id"]))
        by_case[cid].append((line, row))
        for role_name in required:
            if not cell_str(row.get(roles[role_name])):
                findings.append(f"row {line}: required {roles[role_name]!r} is empty")
                critical += 1
        ttype = cell_str(row.get(roles["test_type"]))
        if ttype and ttype not in allowed_types:
            findings.append(f"row {line}: Test Type {ttype!r} not in allowed list")
            critical += 1
        prio = cell_str(row.get(roles["priority"]))
        if prio and prio not in allowed_prio:
            findings.append(f"row {line}: Priority {prio!r} not in allowed list")
            critical += 1
        status = cell_str(row.get(roles["status"]))
        if status and status not in allowed_status:
            findings.append(f"row {line}: Status {status!r} not in allowed list")
            critical += 1
        if status != status_default:
            findings.append(f"row {line}: Status must be {status_default!r} at generation, got {status!r}")
            critical += 1
        if cell_str(row.get(roles["actual_result"])) or cell_str(row.get(roles["defect_id"])):
            findings.append(f"row {line}: Actual Result and Defect ID must be empty at generation")
            critical += 1
        if cid and not id_re.match(cid):
            findings.append(f"row {line}: Test Case ID {cid!r} does not match id pattern")
            critical += 1
        elif cid:
            m = id_re.match(cid)
            scen = cell_str(row.get(roles["scenario_id"]))
            if m and scen.isdigit() and int(m.group(1)) != int(scen):
                findings.append(f"row {line}: Test Case ID prefix {m.group(1)} != Scenario ID {scen}")
                critical += 1

    for cid, items in by_case.items():
        if not cid:
            continue
        if len(items) < min_steps:
            findings.append(f"case {cid}: fewer than {min_steps} step(s)")
            critical += 1
        nums: list[int] = []
        for line, row in items:
            try:
                nums.append(int(cell_str(row[roles["step_number"]])))
            except ValueError:
                findings.append(f"row {line}: Step # is not an integer")
                critical += 1
        if nums and nums != list(range(1, len(nums) + 1)):
            findings.append(f"case {cid}: Step # must be 1..n with no gaps, got {nums}")
            critical += 1
        first = items[0][1]
        for line, row in items[1:]:
            for role_name in case_roles:
                a = cell_str(first[roles[role_name]])
                b = cell_str(row[roles[role_name]])
                if a != b:
                    findings.append(
                        f"row {line}: case-level {roles[role_name]!r} changed within {cid} ({a!r} vs {b!r})"
                    )
                    critical += 1

    return findings, critical
