#!/usr/bin/env python3
"""Regression tests for MSD E2E workbook layout (synthetic fixtures only).

Run: python scripts/test_msd_test_case_xlsx_layout.py

Does not read E2E test execution plan.xlsx or any bootstrap cell content.
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from msd_test_case_xlsx_layout import (  # noqa: E402
    format_test_case_id,
    load_rulebook,
    validate_workbook,
    write_workbook,
)
from importlib.machinery import SourceFileLoader

_cli = SourceFileLoader("msd_test_case_xlsx", str(ROOT / "scripts" / "msd_test_case_xlsx.py")).load_module()
_validate = SourceFileLoader(
    "validate_msd_test_case_xlsx",
    str(ROOT / "scripts" / "validate-msd-test-case-xlsx.py"),
).load_module()


RB_PATH = ROOT / "docs" / "intake" / "msd-test-case-rulebook.yaml"
SRC_RB_PATH = ROOT / "docs" / "intake" / "user-story-csv-rulebook.yaml"


def _cases(rulebook: dict) -> list[dict]:
    return [
        {
            "scenario_id": 1,
            "business_process": "Process Alpha",
            "test_case_id": format_test_case_id(rulebook, 1, "positive", 1),
            "test_case_title": "Complete the alpha capability end to end",
            "test_type": rulebook["test_types"]["positive_default"],
            "priority": rulebook["priority"]["default"],
            "preconditions": "Referenced master data from the user story is available; the cited Dynamics record exists.",
            "source_flow": "synthetic-user-stories.xlsx#US-1",
            "coverage_focus": "Alpha capability happy path",
            "steps": [
                {
                    "step_number": 1,
                    "detailed_test_step": "Open the cited Dynamics app and locate the record named in the story.",
                    "step_owner": "account manager",
                    "application_isv": "Dynamics App A",
                    "test_data": "record cited in the story",
                    "expected_result": "The record is visible and editable.",
                },
                {
                    "step_number": 2,
                    "detailed_test_step": "Perform the capability described in the story need.",
                    "step_owner": "account manager",
                    "application_isv": "Dynamics App A / Dynamics App B",
                    "test_data": "",
                    "expected_result": "The outcome in the story benefit is recorded.",
                },
                {
                    "step_number": 3,
                    "detailed_test_step": "Verify the end-to-end business outcome in the downstream system cited in the story.",
                    "step_owner": "operations",
                    "application_isv": "Dynamics App B",
                    "test_data": "",
                    "expected_result": "The business outcome is visible without manual reconciliation gaps.",
                },
            ],
        },
        {
            "scenario_id": 1,
            "business_process": "Process Alpha",
            "test_case_id": format_test_case_id(rulebook, 1, "exception", 1),
            "test_case_title": "Block alpha capability when the cited constraint fails",
            "test_type": rulebook["test_types"]["exception_default"],
            "priority": rulebook["priority"]["default"],
            "preconditions": "The constraint named in the story AC is not satisfied.",
            "source_flow": "synthetic-user-stories.xlsx#US-1",
            "coverage_focus": "Alpha capability happy path",
            "steps": [
                {
                    "step_number": 1,
                    "detailed_test_step": "Attempt the same capability with the constraint violated.",
                    "step_owner": "account manager",
                    "application_isv": "Dynamics App A",
                    "test_data": "",
                    "expected_result": "The action is refused and the cited control outcome is shown.",
                },
            ],
        },
    ]


def test_good_workbook_passes() -> None:
    rulebook = load_rulebook(RB_PATH)
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        write_workbook(xlsx, rulebook, _cases(rulebook), project="Fixture")
        findings, critical = validate_workbook(rulebook, xlsx)
        assert critical == 0, findings


def test_workbook_applies_rulebook_styling() -> None:
    from openpyxl import load_workbook

    rulebook = load_rulebook(RB_PATH)
    styling = rulebook.get("workbook", {}).get("styling") or {}
    colors = styling.get("colors") or {}
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        write_workbook(xlsx, rulebook, _cases(rulebook), project="Fixture")
        wb = load_workbook(xlsx)
        reg = wb[rulebook["workbook"]["sheets"]["register"]["name"]]
        title_fill = reg.cell(1, 1).fill.fgColor.rgb
        header_fill = reg.cell(4, 1).fill.fgColor.rgb
        assert title_fill == f"FF{colors['title_band']}"
        assert header_fill == f"FF{colors['table_header']}"
        assert reg.cell(1, 1).font.name == styling["font_family"]
        findings, critical = validate_workbook(rulebook, xlsx)
        assert critical == 0, findings


def test_id_pattern_and_denorm() -> None:
    rulebook = load_rulebook(RB_PATH)
    assert format_test_case_id(rulebook, 1, "positive", 1) == "TC01-P01"
    assert format_test_case_id(rulebook, 12, "exception", 1) == "TC12-N01"
    from msd_test_case_xlsx_layout import denormalize_register_rows, load_register_from_xlsx

    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        cases = _cases(rulebook)
        write_workbook(xlsx, rulebook, cases, project="Fixture")
        _headers, rows = load_register_from_xlsx(xlsx, rulebook)
        expected_headers = rulebook["columns"]["register"]["order"]
        assert _headers == expected_headers
        assert len(rows) == 4
        id_col = rulebook["columns"]["register"]["roles"]["test_case_id"]
        title_col = rulebook["columns"]["register"]["roles"]["test_case_title"]
        step_col = rulebook["columns"]["register"]["roles"]["step_number"]
        status_col = rulebook["columns"]["register"]["roles"]["status"]
        actual_col = rulebook["columns"]["register"]["roles"]["actual_result"]
        p_rows = [r for r in rows if r[id_col] == "TC01-P01"]
        assert [r[step_col] for r in p_rows] == ["1", "2", "3"]
        assert len({r[title_col] for r in p_rows}) == 1
        assert all(r[status_col] == rulebook["status_field"]["generation_default"] for r in rows)
        assert all(r[actual_col] == "" for r in rows)
        assert len(denormalize_register_rows(rulebook, cases)) == 4


def test_bad_status_fails() -> None:
    rulebook = load_rulebook(RB_PATH)
    cases = _cases(rulebook)
    cases[0]["steps"][0]["status"] = "Pass"
    # status is forced to generation_default by the writer; mutate the saved file instead
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        write_workbook(xlsx, rulebook, _cases(rulebook), project="Fixture")
        wb = load_workbook(xlsx)
        ws = wb[rulebook["workbook"]["sheets"]["register"]["name"]]
        roles = rulebook["columns"]["register"]["roles"]
        cols = rulebook["columns"]["register"]["order"]
        status_col = cols.index(roles["status"]) + 1
        start = int(rulebook["workbook"]["sheets"]["register"]["data_start_row"])
        ws.cell(start, status_col, "Pass")
        wb.save(xlsx)
        findings, critical = validate_workbook(rulebook, xlsx)
        assert critical > 0
        assert any("Status must be" in f for f in findings)


def test_extract_stories_uses_source_rulebook_headers() -> None:
    from openpyxl import Workbook

    src_rb = load_rulebook(SRC_RB_PATH)
    msd_rb = load_rulebook(RB_PATH)
    cols = list(src_rb["csv"]["column_order"])
    type_col = msd_rb["input"]["row_filter"]["type_column"]
    title_by = ((src_rb.get("columns") or {}).get("stories") or {}).get("title_by_work_item_type") or {}
    feat_title_col = title_by.get("Feature") or "Title 2"
    story_title_col = title_by.get("User Story") or "Title 3"
    if feat_title_col not in cols:
        feat_title_col = msd_rb["input"]["field_bind"]["title"]
    if story_title_col not in cols:
        story_title_col = msd_rb["input"]["field_bind"]["title"]
    desc_col = msd_rb["input"]["field_bind"]["description"]
    ac_col = msd_rb["input"]["field_bind"]["acceptance_criteria"]
    wit_story = src_rb["work_item_types"]["transcript_output_default"]
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "stories.xlsx"
        wb = Workbook()
        ws = wb.active
        for i, h in enumerate(cols, start=1):
            ws.cell(1, i, h)
        # Feature then User Story
        feat = {h: "" for h in cols}
        feat[type_col] = "Feature"
        feat[feat_title_col] = "1. Process Alpha"
        story = {h: "" for h in cols}
        story[type_col] = wit_story
        story[story_title_col] = "Capture the alpha record"
        story[desc_col] = "As an account manager, I need to capture the alpha record, so that downstream work can start."
        story[ac_col] = "- The record is saved.\n- A duplicate is rejected."
        for r_i, data in enumerate((feat, story), start=2):
            for c_i, h in enumerate(cols, start=1):
                ws.cell(r_i, c_i, data[h])
        wb.save(src)
        payload = _cli.extract_stories(src, src_rb, msd_rb)
        assert payload["story_count"] == 1
        assert payload["stories"][0]["nearest_feature_title"] == "1. Process Alpha"
        assert payload["stories"][0]["grouping_key"] == "1. Process Alpha"
        assert payload["stories"][0]["business_process"] == "Process Alpha"


def test_strip_numbering_prefix_is_generic() -> None:
    from msd_test_case_xlsx_layout import strip_numbering_prefix

    rulebook = load_rulebook(RB_PATH)
    assert strip_numbering_prefix(rulebook, "1. Process Alpha") == "Process Alpha"
    assert strip_numbering_prefix(rulebook, "PR-001 – Accrual & Deferral Management") == "Accrual & Deferral Management"
    assert strip_numbering_prefix(rulebook, "#3: Close the period") == "Close the period"
    assert strip_numbering_prefix(rulebook, "Payroll Integration-Deel-Finance") == "Payroll Integration-Deel-Finance"


def test_extract_stories_accepts_title_123_headers() -> None:
    from openpyxl import Workbook

    src_rb = load_rulebook(SRC_RB_PATH)
    msd_rb = load_rulebook(RB_PATH)
    headers = [
        "State",
        "ID",
        "Work Item Type",
        "Title 1",
        "Title 2",
        "Title 3",
        "Area Path",
        "Description",
        "Acceptance Criteria",
    ]
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "workshop.xlsx"
        wb = Workbook()
        ws = wb.active
        for i, h in enumerate(headers, start=1):
            ws.cell(1, i, h)
        ws.cell(2, 2, "90001")
        ws.cell(2, 3, "Feature")
        ws.cell(2, 5, "PR-001 – Process Alpha")
        ws.cell(3, 2, "10001")
        ws.cell(3, 3, "User Story")
        ws.cell(3, 6, "Capture the alpha record")
        ws.cell(3, 7, r"Day 2\Finance")
        ws.cell(3, 8, "As an accountant, I need to capture the alpha record, so that close can start.")
        ws.cell(3, 9, "1. The record is saved.\n2. A duplicate is rejected.")
        wb.save(src)
        payload = _cli.extract_stories(src, src_rb, msd_rb)
        assert payload["story_count"] == 1
        story = payload["stories"][0]
        assert story["title"] == "Capture the alpha record"
        assert story["nearest_feature_title"] == "PR-001 – Process Alpha"
        assert story["business_process"] == "Process Alpha"
        assert story["area_path_leaf"] == "Finance"
        assert "duplicate is rejected" in story["acceptance_criteria"]


def test_writer_cli_roundtrip() -> None:
    rulebook = load_rulebook(RB_PATH)
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        rows_json = Path(tmp) / "rows.json"
        rows_json.write_text(
            json.dumps({"meta": {"project_label": "Fixture"}, "cases": _cases(rulebook)}),
            encoding="utf-8",
        )
        rc = _cli.cmd_write(
            type(
                "A",
                (),
                {
                    "rulebook": str(RB_PATH),
                    "json": str(rows_json),
                    "out": str(xlsx),
                    "project": None,
                    "working_csv": str(Path(tmp) / "working.csv"),
                },
            )()
        )
        assert rc == 0
        findings, critical = validate_workbook(rulebook, xlsx)
        assert critical == 0, findings


def main() -> int:
    tests = [
        test_good_workbook_passes,
        test_workbook_applies_rulebook_styling,
        test_id_pattern_and_denorm,
        test_bad_status_fails,
        test_extract_stories_uses_source_rulebook_headers,
        test_strip_numbering_prefix_is_generic,
        test_extract_stories_accepts_title_123_headers,
        test_writer_cli_roundtrip,
    ]
    failed = 0
    for fn in tests:
        try:
            fn()
            print(f"PASS {fn.__name__}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"FAIL {fn.__name__}: {exc}", file=sys.stderr)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
