#!/usr/bin/env python3
"""Bootstrap MSD365 Business Entity catalog from Area Paths.xlsx (Details sheet only).

Usage:
  python scripts/build-msd365-business-entity-catalog.py
  python scripts/build-msd365-business-entity-catalog.py --xlsx "Area Paths.xlsx" --out docs/intake/msd365-business-entity-catalog.json

Reads columns 2–4 (Domain, Application, Business Entity). Column 1 is ignored.
Column 4 may be empty; empty values are preserved in entries and omitted from business_entities[].
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None  # type: ignore

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "Area Paths.xlsx"
DEFAULT_JSON = ROOT / "docs" / "intake" / "msd365-business-entity-catalog.json"
DEFAULT_CSV = ROOT / "artifacts" / "intake" / ".pipeline" / "msd365-business-entity-catalog.csv"


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def read_details_sheet(xlsx_path: Path) -> list[dict[str, str]]:
    if openpyxl is None:
        print("FAIL: openpyxl required (pip install openpyxl)", file=sys.stderr)
        sys.exit(1)
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if "Details" not in wb.sheetnames:
        print("FAIL: sheet 'Details' not found", file=sys.stderr)
        sys.exit(1)
    ws = wb["Details"]
    entries: list[dict[str, str]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        cells = list(row) + [None, None, None, None]
        domain = _cell(cells[1])
        application = _cell(cells[2])
        business_entity = _cell(cells[3])
        if not domain and not application:
            continue
        entries.append(
            {
                "domain": domain,
                "application": application,
                "business_entity": business_entity,
            }
        )
    return entries


def build_catalog(entries: list[dict[str, str]], source: str) -> dict:
    domains = sorted({e["domain"] for e in entries if e["domain"]})
    applications = sorted({e["application"] for e in entries if e["application"]})
    business_entities = sorted(
        {e["business_entity"] for e in entries if e["business_entity"]}
    )
    return {
        "schema_version": "1.0",
        "catalog_id": "msd365-business-entity-scope-v1",
        "source_bootstrap": source,
        "source_sheet": "Details",
        "column_map": {
            "ignored": "Column 1 (not used for entity scope)",
            "domain": "Column 2 — MSDynamics 365 Domain",
            "application": "Column 3 — MSDynamics 365 Application",
            "business_entity": "Column 4 — MSDynamics 365 Business Entity (may be empty)",
        },
        "entry_count": len(entries),
        "business_entities": business_entities,
        "domains": domains,
        "applications": applications,
        "entries": entries,
    }


def write_csv(entries: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(
            [
                "MSDynamics 365 Domain",
                "MSDynamics 365 Application",
                "MSDynamics 365 Business Entity",
            ]
        )
        for e in entries:
            writer.writerow([e["domain"], e["application"], e["business_entity"]])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--out", type=Path, default=DEFAULT_JSON)
    parser.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"FAIL: xlsx not found: {args.xlsx}", file=sys.stderr)
        sys.exit(1)

    entries = read_details_sheet(args.xlsx)
    catalog = build_catalog(entries, args.xlsx.name)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(catalog, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_csv(entries, args.csv)
    print(
        f"Wrote {args.out} ({catalog['entry_count']} entries, "
        f"{len(catalog['business_entities'])} business entities)"
    )
    print(f"Wrote {args.csv}")


if __name__ == "__main__":
    main()
