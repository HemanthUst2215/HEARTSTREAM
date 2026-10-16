#!/usr/bin/env python3
"""CLI: extract User Story Excel → write MSD E2E workbook from case JSON → validate.

  python scripts/msd_test_case_xlsx.py extract-stories --in stories.xlsx --out stories.json
  python scripts/msd_test_case_xlsx.py write --json cases.json --out artifacts/intake/out.xlsx
  python scripts/msd_test_case_xlsx.py validate --xlsx out.xlsx

Agents: persist case objects to JSON → `write` → `validate`. Never hand-build the workbook.
Column names and sheet layout come only from docs/intake/msd-test-case-rulebook.yaml.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_RULEBOOK = ROOT / "docs" / "intake" / "msd-test-case-rulebook.yaml"
DEFAULT_SOURCE_RULEBOOK = ROOT / "docs" / "intake" / "user-story-csv-rulebook.yaml"

_SCRIPTS = Path(__file__).resolve().parent
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))

from msd_test_case_xlsx_layout import (  # noqa: E402
    area_path_leaf,
    bind_aliases,
    cell_str,
    html_unescape_cell,
    load_rulebook,
    resolve_bound_field,
    resolve_title,
    strip_numbering_prefix,
    write_workbook,
)


def _load_config_project_name(config_path: Path) -> str | None:
    try:
        import yaml  # type: ignore
    except ImportError:
        return None
    if not config_path.is_file():
        return None
    data = yaml.safe_load(config_path.read_text(encoding="utf-8")) or {}
    name = ((data.get("project") or {}).get("name")) if isinstance(data, dict) else None
    return str(name) if name else None


def _read_tabular(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return [[cell_str(c) for c in row] for row in csv.reader(f)]
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("FAIL: openpyxl required (pip install openpyxl)")
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows: list[list[str]] = []
        for r in range(1, ws.max_row + 1):
            rows.append([cell_str(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)])
        return rows
    raise SystemExit(f"FAIL: unsupported input suffix {suffix}")


def _find_header_row(
    grid: list[list[str]],
    msd_rulebook: dict[str, Any],
    source_rulebook: dict[str, Any],
) -> tuple[int, list[str]]:
    """Return (index, actual header cells). Extra columns kept; order is the source row."""
    bind = msd_rulebook["input"]["field_bind"]
    required = list(msd_rulebook["input"].get("required_bind_keys") or ["id", "work_item_type", "title"])
    alias_sets = [bind_aliases(msd_rulebook, key, bind[key]) for key in required]

    expected = list(source_rulebook.get("csv", {}).get("column_order") or [])
    if expected:
        for i, row in enumerate(grid):
            sliced = [cell_str(c) for c in row[: len(expected)]]
            if sliced == expected:
                return i, [cell_str(c) for c in row]
    for i, row in enumerate(grid):
        present = {cell_str(c) for c in row if cell_str(c)}
        if all(any(alias in present for alias in aliases) for aliases in alias_sets):
            return i, [cell_str(c) for c in row]
    needed = "; ".join(f"{k} ({' | '.join(a)})" for k, a in zip(required, alias_sets))
    raise SystemExit(
        "FAIL: could not find a header row with required bind fields: " + needed
    )


def _row_dict(header: list[str], cells: list[str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for i, h in enumerate(header):
        name = cell_str(h)
        if not name or name in out:
            continue
        out[name] = cells[i] if i < len(cells) else ""
    return out


def extract_stories(source: Path, source_rulebook: dict[str, Any], msd_rulebook: dict[str, Any]) -> dict[str, Any]:
    grid = _read_tabular(source)
    unescape = bool(((msd_rulebook.get("input") or {}).get("cell_normalize") or {}).get("html_unescape"))
    if unescape:
        grid = [[html_unescape_cell(c, True) for c in row] for row in grid]
    header_idx, header = _find_header_row(grid, msd_rulebook, source_rulebook)
    type_col = msd_rulebook["input"]["row_filter"]["type_column"]
    include_key = msd_rulebook["input"]["row_filter"]["include_type_from_source"]
    include_type = source_rulebook.get("work_item_types", {}).get(include_key) or "User Story"
    feature_type = msd_rulebook["input"]["grouping"]["nearest_preceding_feature"]["type_value"]
    bind = msd_rulebook["input"]["field_bind"]

    last_feature = {"title": "", "id": "", "row": None}
    last_epic = {"title": "", "id": "", "row": None}
    stories: list[dict[str, Any]] = []
    context_rows: list[dict[str, Any]] = []

    def _bound(row: dict[str, str], wit: str) -> dict[str, str]:
        return {
            "id": resolve_bound_field(row, msd_rulebook, "id", bind["id"]),
            "title": resolve_title(row, wit, msd_rulebook, bind["title"]),
            "description": resolve_bound_field(row, msd_rulebook, "description", bind["description"]),
            "acceptance_criteria": resolve_bound_field(
                row, msd_rulebook, "acceptance_criteria", bind["acceptance_criteria"]
            ),
            "area_path": resolve_bound_field(row, msd_rulebook, "area_path", bind["area_path"]),
        }

    for offset, cells in enumerate(grid[header_idx + 1 :], start=header_idx + 2):
        if all(c == "" for c in cells):
            continue
        row = _row_dict(header, cells)
        wit = cell_str(row.get(type_col))
        fields = _bound(row, wit)
        if wit == "Epic":
            last_epic = {"title": fields["title"], "id": fields["id"], "row": offset}
            context_rows.append({"work_item_type": wit, "row": offset, **last_epic})
            continue
        if wit == feature_type:
            last_feature = {"title": fields["title"], "id": fields["id"], "row": offset}
            context_rows.append({"work_item_type": wit, "row": offset, **last_feature})
            continue
        if wit != include_type:
            continue
        stories.append(
            {
                "row": offset,
                "id": fields["id"],
                "work_item_type": wit,
                "title": fields["title"],
                "description": fields["description"],
                "acceptance_criteria": fields["acceptance_criteria"],
                "area_path": fields["area_path"],
                "area_path_leaf": area_path_leaf(fields["area_path"]),
                "nearest_feature_title": last_feature["title"],
                "nearest_feature_id": last_feature["id"],
                "nearest_epic_title": last_epic["title"],
                "nearest_epic_id": last_epic["id"],
                "source_path": str(source),
                "source_locator": f"{source.name}#row-{offset}",
            }
        )

    features_by_id = {
        cell_str(r.get("id")): r
        for r in context_rows
        if r.get("work_item_type") == feature_type and cell_str(r.get("id"))
    }
    grouping_key = []
    for s in stories:
        mapped = features_by_id.get(s["id"])
        if mapped:
            s["nearest_feature_title"] = mapped["title"]
            s["nearest_feature_id"] = mapped["id"]
            key = mapped["title"]
        else:
            key = s["nearest_feature_title"] or s["area_path"] or s["title"]
        grouping_key.append(key)
        s["grouping_key"] = key
        s["business_process"] = strip_numbering_prefix(msd_rulebook, key)

    return {
        "source": str(source),
        "header_row": header_idx + 1,
        "include_type": include_type,
        "story_count": len(stories),
        "group_count": len({k for k in grouping_key if k}),
        "stories": stories,
        "context_rows": context_rows,
    }


def cmd_extract(args: argparse.Namespace) -> int:
    msd = load_rulebook(Path(args.rulebook))
    source_rb = load_rulebook(Path(args.source_rulebook))
    payload = extract_stories(Path(args.input), source_rb, msd)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {out} ({payload['story_count']} stories, {payload['group_count']} groups)")
    return 0


def cmd_write(args: argparse.Namespace) -> int:
    rulebook = load_rulebook(Path(args.rulebook))
    payload = json.loads(Path(args.json).read_text(encoding="utf-8"))
    if isinstance(payload, list):
        cases = payload
        meta: dict[str, Any] = {}
    elif isinstance(payload, dict):
        cases = payload.get("cases")
        if not isinstance(cases, list):
            raise SystemExit("JSON must contain a 'cases' array")
        meta = payload.get("meta") or {}
    else:
        raise SystemExit("JSON root must be an object with cases[] or an array")
    project = args.project or meta.get("project_label") or _load_config_project_name(ROOT / ".cursor" / "config.yaml")
    csv_path = Path(args.working_csv) if args.working_csv else None
    result = write_workbook(Path(args.out), rulebook, cases, project=project, also_csv=csv_path)
    print(json.dumps(result, indent=2))
    return 0


def cmd_validate(args: argparse.Namespace) -> int:
    import subprocess

    cmd = [
        sys.executable,
        str(ROOT / "scripts" / "validate-msd-test-case-xlsx.py"),
        "--xlsx",
        args.xlsx,
        "--rulebook",
        args.rulebook,
    ]
    if args.report:
        cmd.extend(["--report", args.report])
    return subprocess.call(cmd)


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)

    e = sub.add_parser("extract-stories", help="Read User Story Excel/CSV using source rulebook headers")
    e.add_argument("--in", dest="input", required=True)
    e.add_argument("--out", required=True)
    e.add_argument("--rulebook", default=str(DEFAULT_RULEBOOK))
    e.add_argument("--source-rulebook", default=str(DEFAULT_SOURCE_RULEBOOK))

    w = sub.add_parser("write", help="Write MSD E2E xlsx from case JSON")
    w.add_argument("--json", required=True)
    w.add_argument("--out", required=True)
    w.add_argument("--rulebook", default=str(DEFAULT_RULEBOOK))
    w.add_argument("--project")
    w.add_argument("--working-csv", help="Optional internal register CSV")

    v = sub.add_parser("validate", help="Validate generated workbook")
    v.add_argument("--xlsx", required=True)
    v.add_argument("--rulebook", default=str(DEFAULT_RULEBOOK))
    v.add_argument("--report")

    args = parser.parse_args()
    if args.command == "extract-stories":
        return cmd_extract(args)
    if args.command == "write":
        return cmd_write(args)
    return cmd_validate(args)


if __name__ == "__main__":
    sys.exit(main())
