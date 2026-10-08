#!/usr/bin/env python3
"""Convert a markdown file to a Word (.docx) document.

Focused converter for the Copilot vs Cursor comparison doc.
Handles: H1-H3 headings, pipe tables, bullet/ordered lists, blockquotes,
horizontal rules, bold/italic/code inline, and links.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


INLINE_RE = re.compile(
    r"(\*\*(?P<bold>[^*]+)\*\*)"
    r"|(\*(?P<italic>[^*]+)\*)"
    r"|(`(?P<code>[^`]+)`)"
    r"|(\[(?P<text>[^\]]+)\]\((?P<url>[^)]+)\))"
)


def set_cell_shading(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def add_inline(paragraph, text):
    pos = 0
    while pos < len(text):
        m = INLINE_RE.search(text, pos)
        if not m:
            paragraph.add_run(text[pos:])
            break
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        if m.group("bold"):
            r = paragraph.add_run(m.group("bold"))
            r.bold = True
        elif m.group("italic"):
            r = paragraph.add_run(m.group("italic"))
            r.italic = True
        elif m.group("code"):
            r = paragraph.add_run(m.group("code"))
            r.font.name = "Menlo"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
        elif m.group("text"):
            r = paragraph.add_run(m.group("text"))
            r.font.color.rgb = RGBColor(0x0B, 0x57, 0xD0)
            r.underline = True
        pos = m.end()


def parse_table(lines, idx):
    rows = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        rows.append(lines[idx].strip())
        idx += 1
    # drop separator row
    if len(rows) >= 2 and re.match(r"^\|[\s:|-]+\|$", rows[1].replace("|", "|")):
        del rows[1]
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip("|").split("|")]
        parsed.append(cells)
    return parsed, idx


def add_table(doc, rows):
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        for j in range(n_cols):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, row[j] if j < len(row) else "")
            if i == 0:
                for run in para.runs:
                    run.bold = True
                set_cell_shading(cell, "D9E2F3")
    doc.add_paragraph("")


def convert(md_path, docx_path):
    text = Path(md_path).read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    # base style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    in_list = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # blank line
        if not stripped:
            in_list = False
            i += 1
            continue

        # horizontal rule
        if stripped in ("---", "***", "___"):
            p = doc.add_paragraph()
            p_pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)
            p_pr.append(pbdr)
            i += 1
            continue

        # headings
        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if hm:
            level = len(hm.group(1))
            title = hm.group(2).strip()
            h = doc.add_heading(level=min(level, 3))
            add_inline(h, title)
            i += 1
            continue

        # blockquote
        if stripped.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, "\n".join(quote_lines))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        # table
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            in_list = False
            continue

        # ordered list
        om = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if om:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, om.group(2))
            i += 1
            continue

        # unordered list
        if stripped.startswith("- ") or stripped.startswith("* "):
            content = stripped[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, content)
            i += 1
            continue

        # sub-bullet
        if re.match(r"^\s+[-*]\s+", line):
            content = re.sub(r"^\s+[-*]\s+", "", line)
            p = doc.add_paragraph(style="List Bullet 2")
            add_inline(p, content)
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"Wrote {docx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: md_to_docx.py <input.md> <output.docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
