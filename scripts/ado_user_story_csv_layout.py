#!/usr/bin/env python3
"""Canonical ADO user-story CSV/xlsx layout: write helpers + structural validation.

All `/generate-user-stories-from-transcript` CSV output must be written through
`write_csv()` and pass `validate_structure()` (also invoked by validate-user-story-csv.py).
Public Excel styling comes from docs/intake/user-story-csv-rulebook.yaml workbook.styling.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore
    Alignment = None  # type: ignore
    Font = None  # type: ignore
    PatternFill = None  # type: ignore
    get_column_letter = None  # type: ignore

USER_STORY_RE = re.compile(
    r"^As an?\s+.+\s*,?\s*I need to\s+.+\s*,?\s*so that\s+.+",
    re.IGNORECASE | re.DOTALL,
)

# Monotonic section order (rulebook ado_export_layout.row_order)
SECTION_RANK: dict[str, int] = {
    "Workshop": 0,
    "Epic": 1,
    "Feature": 2,
    "Test Case": 3,
    "User Story": 3,
    "Task": 4,
    "Issue": 5,
    "Risk": 6,
}

TEST_CASE_HEAD_RE = re.compile(
    r"^Transcript:\s*\S.+\n"
    r"Epic:\s*(?P<epic>\d+)\s*\|\s*\S.+\n"
    r"Feature:\s*(?P<feature>\d+)\s*\|\s*\S.+\n"
    r"Steps:\n"
    r"(?:\d+\.\s+\S.*\n?)+",
    re.MULTILINE,
)

# Transcript authoring does not create Epic/Feature rows; those IDs come from work-items JSON.
HIERARCHY_BEFORE_STORIES = frozenset({"Workshop"})

# ASCII hyphen bullets (Excel/ADO-friendly); do not use Unicode U+2022 (•).
AC_BULLET_PREFIX = "- "


DEFAULT_TITLE_BY_TYPE: dict[str, str] = {
    "Workshop": "Title 1",
    "Epic": "Title 1",
    "Feature": "Title 2",
    "User Story": "Title 3",
    "Test Case": "Title 3",
    "Task": "Title 3",
    "Issue": "Title 3",
    "Risk": "Title 3",
}

DEFAULT_TITLE_RESOLUTION: dict[str, list[str]] = {
    "Workshop": ["Title 1", "Title"],
    "Epic": ["Title 1", "Title"],
    "Feature": ["Title 2", "Title"],
    "User Story": ["Title 3", "Title"],
    "Test Case": ["Title 3", "Title"],
    "Task": ["Title 3", "Title"],
    "Issue": ["Title 3", "Title"],
    "Risk": ["Title 3", "Title"],
}


def headers_from_rulebook(rulebook: dict[str, Any]) -> list[str]:
    cols = (
        ((rulebook.get("columns") or {}).get("stories") or {}).get("order")
        or (rulebook.get("csv") or {}).get("column_order")
    )
    if not cols:
        raise ValueError("rulebook columns.stories.order / csv.column_order missing")
    return list(cols)


def title_column_for_type(rulebook: dict[str, Any], wit: str) -> str:
    mapping = ((rulebook.get("columns") or {}).get("stories") or {}).get(
        "title_by_work_item_type"
    ) or {}
    return str(mapping.get(wit) or DEFAULT_TITLE_BY_TYPE.get(wit) or "Title 3")


def _title_candidates(rulebook: dict[str, Any], wit: str) -> list[str]:
    reso = (rulebook.get("input") or {}).get("title_resolution") or {}
    by_type = reso.get("by_work_item_type") or {}
    default = list(reso.get("default") or ["Title 3", "Title", "Title 2", "Title 1"])
    return list(by_type.get(wit) or DEFAULT_TITLE_RESOLUTION.get(wit) or default)


def row_title(row: dict[str, str], rulebook: dict[str, Any] | None = None, wit: str | None = None) -> str:
    """First non-empty title for this work item type (Title 1/2/3 or legacy Title)."""
    rb = rulebook or {}
    kind = (wit or row.get("Work Item Type") or "").strip()
    for col in _title_candidates(rb, kind):
        val = (row.get(col) or "").strip()
        if val:
            return val
    return ""


def apply_title_columns(rulebook: dict[str, Any], row: dict[str, str]) -> dict[str, str]:
    """Copy a legacy `Title` (or first resolved title) into the type-specific Title 1/2/3 cell."""
    wit = (row.get("Work Item Type") or "").strip()
    dest = title_column_for_type(rulebook, wit)
    if dest and not (row.get(dest) or "").strip():
        title = row_title(row, rulebook, wit)
        if title:
            row[dest] = title
    return row


def apply_generation_defaults(rulebook: dict[str, Any], row: dict[str, str]) -> dict[str, str]:
    """Fill State default when the column exists and the cell is empty. Does not invent owners."""
    wit = (row.get("Work Item Type") or "").strip()
    state_spec = (rulebook.get("generation") or {}).get("state_field") or {}
    state_col = str(state_spec.get("column") or "State")
    defaults = state_spec.get("defaults") or {}
    if state_col in row and not (row.get(state_col) or "").strip():
        default = defaults.get(wit)
        if default:
            row[state_col] = str(default)
    apply_title_columns(rulebook, row)
    return row


def blank_row(headers: list[str]) -> dict[str, str]:
    return {h: "" for h in headers}


def ac_bullet_lines(ac: str) -> list[str]:
    return [ln.strip() for ln in ac.splitlines() if ln.strip().startswith(AC_BULLET_PREFIX)]


def acceptance_criteria_cell(*bullets: str) -> str:
    """One ASCII `- ` bullet per line (required when len(bullets) > 1)."""
    return "\n".join(f"{AC_BULLET_PREFIX}{b.strip()}" for b in bullets if b.strip())


def issue_description_cell(risk: str, impact: str, mitigation: str) -> str:
    return f"Risk: {risk.strip()}\nImpact: {impact.strip()}\nMitigation: {mitigation.strip()}"


# Characters that Excel (and some other spreadsheet tools) treat as a formula
# trigger when they are the first character of a CSV cell. A cell like
# "- Partner record is not created..." gets parsed as `-Partner` minus
# something and renders as `#NAME?` once opened in Excel. RFC4180 quoting
# alone does not prevent this -- Excel re-applies its own formula detection
# to the *unquoted* cell content after parsing the CSV. Prefixing such cells
# with a single space neutralizes the trigger while staying fully readable
# (the validator strips leading whitespace before checking the "- " bullet
# prefix, and ADO/plain-text consumers render the content unchanged).
EXCEL_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def excel_safe_cell(value: str) -> str:
    """Prefix a cell value with a single space if it would be mis-parsed as
    an Excel formula (leading =, +, -, or @). No-op for empty/other values.
    """
    if value and value[0] in EXCEL_FORMULA_TRIGGER_CHARS:
        return f" {value}"
    return value


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    """RFC4180 CSV with quoted fields that contain newlines or commas.

    Cell values are passed through `excel_safe_cell()` so that bullet lists
    (Acceptance Criteria `- ` bullets) and any other field beginning with a
    formula-trigger character do not render as `#NAME?` when opened in Excel.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    for row in rows:
        missing = set(headers) - set(row.keys())
        if missing:
            raise ValueError(f"row missing columns: {sorted(missing)}")
    safe_rows = [{k: excel_safe_cell(v) for k, v in row.items()} for row in rows]
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=headers,
            lineterminator="\n",
            quoting=csv.QUOTE_MINIMAL,
        )
        writer.writeheader()
        writer.writerows(safe_rows)


def validate_structure(
    rulebook: dict[str, Any],
    rows: list[dict[str, str]],
    row_offset: int = 2,
) -> tuple[list[str], int]:
    """Return (findings, critical_count)."""
    findings: list[str] = []
    critical = 0

    layout = rulebook.get("ado_export_layout", {})
    us_area_default = (
        rulebook.get("required_fields", {})
        .get("User Story", {})
        .get("area_path_default", "optional")
    )

    types_seen: set[str] = set()
    prev_rank = -1
    story_count = 0
    saw_story = False
    saw_test_case = False
    epic_ids: set[str] = set()
    feature_ids: set[str] = set()
    cited_pairs: list[tuple[int, str, str]] = []

    for i, row in enumerate(rows):
        line = row_offset + i
        apply_generation_defaults(rulebook, row)
        wit = (row.get("Work Item Type") or "").strip()
        if not wit:
            findings.append(f"row {line}: missing Work Item Type")
            critical += 1
            continue

        rank = SECTION_RANK.get(wit, 99)
        if rank < prev_rank:
            findings.append(
                f"row {line}: {wit} is out of ADO section order "
                f"(expected Workshop → User Story → Task → Issue → Risk; Epic/Feature are catalog IDs, not new rows)"
            )
            critical += 1
        prev_rank = max(prev_rank, rank)

        if wit == "User Story":
            story_count += 1
            if not saw_story:
                saw_story = True
                if not HIERARCHY_BEFORE_STORIES.issubset(types_seen):
                    missing = sorted(HIERARCHY_BEFORE_STORIES - types_seen)
                    findings.append(
                        f"row {line}: first User Story before hierarchy rows present: missing {missing}"
                    )
                    critical += 1
            area = (row.get("Area Path") or "").strip()
            if us_area_default == "empty" and area:
                findings.append(
                    f"row {line}: User Story Area Path must be empty (ADO export convention)"
                )
                critical += 1
            ac = row.get("Acceptance Criteria") or ""
            if ac and "\u2022" in ac:
                findings.append(
                    f"row {line}: Acceptance Criteria must use ASCII '- ' bullets, not Unicode • (U+2022)"
                )
                critical += 1
            bullet_lines = ac_bullet_lines(ac)
            if ac and not bullet_lines:
                findings.append(
                    f"row {line}: Acceptance Criteria must use lines starting with '- ' (hyphen + space)"
                )
                critical += 1
            else:
                for ac_line in ac.splitlines():
                    if ac_line.count(AC_BULLET_PREFIX) > 1:
                        findings.append(
                            f"row {line}: Acceptance Criteria must use one '- ' bullet per line, not multiple on one line"
                        )
                        critical += 1
                        break
                if len(bullet_lines) >= 2 and len(ac.splitlines()) < 2:
                    findings.append(
                        f"row {line}: Acceptance Criteria must use one '- ' bullet per line, not a single line"
                    )
                    critical += 1

        if wit in ("Issue", "Risk"):
            ac = (row.get("Acceptance Criteria") or "").strip()
            if ac:
                findings.append(f"row {line}: {wit} must have empty Acceptance Criteria")
                critical += 1
            desc = row.get("Description") or ""
            for label in ("Risk:", "Impact:", "Mitigation:"):
                if label not in desc:
                    findings.append(f"row {line}: {wit} Description must include {label}")
                    critical += 1

        if wit == "Workshop":
            area = (row.get("Area Path") or "").strip()
            if layout.get("workshop_row") and not area:
                findings.append(f"row {line}: Workshop should include Area Path when known")
                # warning only — not critical
            annex = (row.get("Annexure ID1") or "").strip()
            if annex:
                findings.append(f"row {line}: Workshop must have empty Annexure ID1 (annexure belongs on SOW Feature rows only)")
                critical += 1

        if wit == "Epic":
            epic_id = (row.get("ID") or "").strip()
            if epic_id:
                epic_ids.add(epic_id)
        if wit == "Feature":
            feature_id = (row.get("ID") or "").strip()
            if feature_id:
                feature_ids.add(feature_id)

        if wit == "Test Case":
            if not saw_test_case:
                saw_test_case = True
                if "Epic" not in types_seen or "Feature" not in types_seen:
                    findings.append(
                        f"row {line}: Test Case before Epic and Feature catalog rows"
                    )
                    critical += 1
            if (row.get("ID") or "").strip():
                findings.append(f"row {line}: Test Case ID must be blank")
                critical += 1
            desc = row.get("Description") or ""
            match = TEST_CASE_HEAD_RE.search(desc.strip())
            if not match:
                findings.append(
                    "row {line}: Test Case Description must be "
                    "Transcript / Epic: <id> | <title> / Feature: <id> | <title> / Steps".format(
                        line=line
                    )
                )
                critical += 1
            else:
                cited_pairs.append((line, match.group("epic"), match.group("feature")))
            ac = row.get("Acceptance Criteria") or ""
            if ac and "\u2022" in ac:
                findings.append(
                    f"row {line}: Acceptance Criteria must use ASCII '- ' bullets, not Unicode • (U+2022)"
                )
                critical += 1
            if not ac_bullet_lines(ac):
                findings.append(
                    f"row {line}: Test Case Acceptance Criteria must use lines starting with '- '"
                )
                critical += 1

        types_seen.add(wit)

    if story_count and "Workshop" not in types_seen:
        findings.append("file: at least one Workshop row required when User Stories are present")
        critical += 1

    if saw_test_case:
        for line, epic_id, feature_id in cited_pairs:
            if epic_id not in epic_ids:
                findings.append(
                    f"row {line}: Test Case Epic {epic_id} is not an Epic row ID in this file"
                )
                critical += 1
            if feature_id not in feature_ids:
                findings.append(
                    f"row {line}: Test Case Feature {feature_id} is not a Feature row ID in this file"
                )
                critical += 1

    return findings, critical


def validate_row_content(
    rulebook: dict[str, Any],
    rows: list[dict[str, str]],
    row_offset: int = 2,
) -> tuple[list[str], int]:
    """Field-level rules (narrative, AC presence)."""
    findings: list[str] = []
    critical = 0
    required_by_type = rulebook.get("required_fields", {})
    allowed_types = set(rulebook.get("work_item_types", {}).get("allowed", []))

    for i, row in enumerate(rows):
        line = row_offset + i
        apply_generation_defaults(rulebook, row)
        wit = (row.get("Work Item Type") or "").strip()
        if wit not in allowed_types:
            continue
        for col in required_by_type.get(wit, {}).get("required", []):
            if col == "workshop_governance_columns":
                continue
            if not (row.get(col) or "").strip():
                findings.append(f"row {line} ({wit}): required column '{col}' is empty")
                critical += 1
        if wit == "User Story":
            desc = row.get("Description") or ""
            if not USER_STORY_RE.search(desc):
                findings.append(
                    f"row {line}: User Story Description must match "
                    "'As a <role>, I need to ..., so that ...'"
                )
                critical += 1
    return findings, critical


def _argb(hex_color: str) -> str:
    raw = (hex_color or "").strip().lstrip("#")
    if len(raw) == 6:
        return "FF" + raw.upper()
    if len(raw) == 8:
        return raw.upper()
    return "FFFFFFFF"


def stories_sheet_name(rulebook: dict[str, Any], project_label: str | None = None) -> str:
    sheet = ((rulebook.get("workbook") or {}).get("sheets") or {}).get("stories") or {}
    spec = (rulebook.get("workbook") or {}).get("project_label") or {}
    placeholders = set(spec.get("placeholder_tokens") or ["<project-name>", ""])
    fallback = str(sheet.get("fallback_name") or spec.get("fallback") or "User Stories- Workshop")
    label = (project_label or "").strip()
    if not label or label in placeholders:
        name = fallback
    else:
        tmpl = str(sheet.get("name_template") or "{project_label} User Stories- Workshop")
        name = tmpl.format(project_label=label)
    max_len = int(sheet.get("max_name_length") or 31)
    cleaned = re.sub(r'[:\\/?*\[\]]', " ", name).strip() or fallback
    return cleaned[:max_len]


def write_xlsx(
    path: Path,
    headers: list[str],
    rows: list[dict[str, str]],
    rulebook: dict[str, Any] | None = None,
    project_label: str | None = None,
) -> None:
    """Public User Story workbook with rulebook styling. Requires openpyxl."""
    if Workbook is None:
        raise SystemExit("FAIL: openpyxl required (pip install openpyxl)")
    rb = rulebook or {}
    styling = (rb.get("workbook") or {}).get("styling") or {}
    sheet_style = ((styling.get("sheets") or {}).get("stories") or {})
    colors = styling.get("colors") or {}
    font_name = str(styling.get("font_family") or "Aptos Narrow")
    body_size = float(styling.get("body_font_size") or 9)
    header_size = float(styling.get("header_font_size") or body_size)
    header_bold = bool(styling.get("header_bold", False))
    header_height = float(styling.get("header_row_height") or 24)
    wrap = bool(styling.get("wrap_text", True))
    vertical = str(styling.get("vertical") or "top")
    widths = sheet_style.get("column_widths_by_header") or {}
    hidden = set(sheet_style.get("hidden_columns") or [])
    highlight = sheet_style.get("highlight") or {}
    highlight_type = str(highlight.get("when_work_item_type") or "Workshop")
    highlight_cols = set(highlight.get("fill_columns") or ["Work Item Type", "Title 1"])
    fill_key = str(highlight.get("fill") or "workshop_highlight")
    highlight_hex = str(colors.get(fill_key) or "FFFF00")
    auto_h = sheet_style.get("auto_row_height") or {}
    min_h = float(auto_h.get("min") or 15)
    max_h = float(auto_h.get("max") or 280)
    line_h = float(auto_h.get("line_height") or 12)

    body_font = Font(name=font_name, size=body_size, bold=False)
    header_font = Font(name=font_name, size=header_size, bold=header_bold)
    align = Alignment(wrap_text=wrap, vertical=vertical)
    workshop_fill = PatternFill("solid", fgColor=_argb(highlight_hex))

    wb = Workbook()
    ws = wb.active
    ws.title = stories_sheet_name(rb, project_label)

    for c, header in enumerate(headers, start=1):
        cell = ws.cell(1, c, header)
        cell.font = header_font
        cell.alignment = align
    ws.row_dimensions[1].height = header_height

    for r_i, row in enumerate(rows, start=2):
        apply_generation_defaults(rb, row)
        wit = (row.get("Work Item Type") or "").strip()
        longest = 1
        for c, header in enumerate(headers, start=1):
            value = excel_safe_cell(row.get(header) or "")
            cell = ws.cell(r_i, c, value)
            cell.font = body_font
            cell.alignment = align
            if wit == highlight_type and header in highlight_cols:
                cell.fill = workshop_fill
            if value:
                longest = max(longest, value.count("\n") + 1)
        ws.row_dimensions[r_i].height = min(max_h, max(min_h, longest * line_h))

    for c, header in enumerate(headers, start=1):
        letter = get_column_letter(c)
        if header in hidden:
            ws.column_dimensions[letter].hidden = True
            ws.column_dimensions[letter].width = float(widths.get(header) or 0)
        elif header in widths:
            ws.column_dimensions[letter].width = float(widths[header])

    freeze = ((rb.get("workbook") or {}).get("sheets") or {}).get("stories", {}).get("freeze_panes")
    if freeze:
        ws.freeze_panes = freeze

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)

